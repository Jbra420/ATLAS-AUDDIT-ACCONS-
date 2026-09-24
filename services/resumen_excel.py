"""
services/resumen_excel.py — Levantamiento de información del expediente en Excel (.xlsx).

Reúne en un libro todo lo que pide el levantamiento de información: los
bloques 1 a 6 (SRI, Supercias, ubicación, administradores, accionistas e
información financiera) con la fuente y la fecha de consulta de cada dato, el
estado de los requisitos, las validaciones cruzadas y alertas, las fuentes y la
evidencia, la trazabilidad, los hallazgos y la recomendación preliminar.

Lee el expediente tal como está: no genera ni guarda el resumen. Usa las
mismas reglas que el resumen en texto (services/summary.py y
services/validaciones.py), así que ambos coinciden.

Los valores se escriben siempre como datos: un texto que empieza con "=" (por
ejemplo, pegado desde una web) se guarda como texto y nunca como fórmula.
"""
from __future__ import annotations

import io
from datetime import datetime
from typing import Any, Iterable

from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from services.financial import compute_indicators, filas_comparativo
from services.rowutil import row_get
from services.summary import (
    CAMPOS_SRI,
    CAMPOS_UBICACION,
    _CamposTolerantes,
    cierre_levantamiento,
    extract_signals,
    risk_suggestions,
    valores_supercias,
)
from services.trazabilidad import (
    BLOQUE_ACCIONISTAS,
    BLOQUE_ADMINISTRADORES,
    BLOQUE_FINANCIERO,
    BLOQUE_SRI,
    BLOQUE_SUPERCIAS,
    BLOQUE_UBICACION,
)
from services.validaciones import evaluar_levantamiento

PENDIENTE = "Pendiente de confirmar"
AVISO = (
    "Documento PRELIMINAR generado con la información registrada en Atlas. Debe validarse contra "
    "fuentes oficiales (SRI, Supercias, SERCOP) antes de utilizarse como soporte de auditoría formal."
)

_FORMATO_USD = '"$"#,##0.00;[Red]-"$"#,##0.00'
_FORMATO_PCT = "0.00%"
_ESTADOS_CRUCE = {
    "coincide": "Coincide", "no_coincide": "No coincide",
    "pendiente": "Pendiente de datos", "revisar": "Revisar",
}
_NIVELES = {"critica": "Crítica", "alta": "Alta", "media": "Media", "informativa": "Informativa"}
_BLOQUES = {
    BLOQUE_SRI: "SRI", BLOQUE_SUPERCIAS: "Supercias", BLOQUE_UBICACION: "Ubicación",
    BLOQUE_ADMINISTRADORES: "Administradores", BLOQUE_ACCIONISTAS: "Accionistas",
    BLOQUE_FINANCIERO: "Financiero",
}

# Estilo del libro: marca Atlas, sobrio y legible impreso.
_MORADO = "5B2A86"
_TITULO = Font(bold=True, size=14, color=_MORADO)
_SUBTITULO = Font(bold=True, size=11, color=_MORADO)
_ENCABEZADO = Font(bold=True, color="FFFFFF")
_RELLENO_ENCABEZADO = PatternFill("solid", fgColor=_MORADO)
_RELLENO_TOTAL = PatternFill("solid", fgColor="EEE8F5")
_RELLENO_PENDIENTE = PatternFill("solid", fgColor="FFF4E5")
_RELLENO_ALERTA = PatternFill("solid", fgColor="FDECEC")
_BORDE = Border(bottom=Side(style="thin", color="D9D2E3"))
_AJUSTE = Alignment(wrap_text=True, vertical="top")


# ---------------------------------------------------------------------------
# Escritura segura
# ---------------------------------------------------------------------------

def _texto(value: Any) -> Any:
    """Texto sin caracteres de control (openpyxl los rechaza)."""
    return ILLEGAL_CHARACTERS_RE.sub("", value) if isinstance(value, str) else value


def _poner(ws: Worksheet, fila: int, columna: int, value: Any, formato: str = "") -> Any:
    """Escribe un valor como dato. openpyxl interpreta como fórmula todo texto
    que empieza con "=": se fuerza a texto para evitar inyección de fórmulas."""
    cell = ws.cell(row=fila, column=columna, value=_texto(value))
    if isinstance(cell.value, str) and cell.value.startswith("="):
        cell.data_type = "s"
    if formato:
        cell.number_format = formato
    cell.alignment = _AJUSTE
    return cell


def _vacio(value: Any) -> bool:
    return value is None or (isinstance(value, str) and value.strip() in {"", "-", "—"})


def _o_pendiente(value: Any) -> Any:
    return PENDIENTE if _vacio(value) else value


class _Hoja:
    """Hoja que se llena de arriba hacia abajo: título, pares campo/valor y tablas."""

    def __init__(self, libro: Workbook, titulo: str, anchos: Iterable[int]) -> None:
        self.ws = libro.create_sheet(titulo[:31])
        self.fila = 1
        for i, ancho in enumerate(anchos, start=1):
            self.ws.column_dimensions[get_column_letter(i)].width = ancho
        self.ws.sheet_view.showGridLines = False
        self.ws.page_setup.orientation = "landscape"
        self.ws.page_setup.fitToWidth = 1
        self.ws.page_setup.fitToHeight = 0
        self.ws.sheet_properties.pageSetUpPr.fitToPage = True

    def titulo(self, texto: str, subtitulo: str = "") -> None:
        _poner(self.ws, self.fila, 1, texto).font = _TITULO
        self.fila += 1
        if subtitulo:
            _poner(self.ws, self.fila, 1, subtitulo).font = Font(italic=True, color="666666")
            self.fila += 1
        self.fila += 1

    def seccion(self, texto: str) -> None:
        _poner(self.ws, self.fila, 1, texto).font = _SUBTITULO
        self.fila += 1

    def tabla(
        self, encabezados: list[str], filas: Iterable[Iterable[Any]], formatos: dict[int, str] | None = None,
        vacia: str = "Sin registros.", rellenos: Iterable[PatternFill | None] | None = None,
        fijar_encabezado: bool = False,
    ) -> tuple[int, int]:
        """Escribe una tabla y devuelve (primera, última) fila de datos.
        fijar_encabezado deja visible la fila de títulos al desplazarse."""
        formatos = formatos or {}
        for col, encabezado in enumerate(encabezados, start=1):
            cell = _poner(self.ws, self.fila, col, encabezado)
            cell.font, cell.fill = _ENCABEZADO, _RELLENO_ENCABEZADO
        if fijar_encabezado:
            self.ws.freeze_panes = self.ws.cell(row=self.fila + 1, column=1)
        self.fila += 1
        primera = self.fila
        rellenos = list(rellenos or [])
        filas = list(filas)
        for i, fila in enumerate(filas):
            for col, value in enumerate(fila, start=1):
                cell = _poner(self.ws, self.fila, col, value, formatos.get(col, ""))
                cell.border = _BORDE
                if i < len(rellenos) and rellenos[i] is not None:
                    cell.fill = rellenos[i]
            self.fila += 1
        if not filas:
            _poner(self.ws, self.fila, 1, vacia).font = Font(italic=True, color="666666")
            self.fila += 1
        ultima = self.fila - 1
        self.fila += 1
        return primera, ultima

    def espacio(self) -> None:
        self.fila += 1


# ---------------------------------------------------------------------------
# Datos del expediente
# ---------------------------------------------------------------------------

def _trazabilidad(provenance: list) -> dict[tuple[str, str], Any]:
    """Último registro de trazabilidad de cada (bloque, campo)."""
    ultimo: dict[tuple[str, str], Any] = {}
    for row in provenance or []:
        ultimo[(row_get(row, "bloque"), row_get(row, "campo"))] = row
    return ultimo


def _fuente(traza: dict, bloque: str, campo: str) -> tuple[str, str]:
    row = traza.get((bloque, campo))
    if row is None:
        return "Sin trazabilidad registrada", ""
    return row_get(row, "fuente") or "", row_get(row, "fecha_consulta") or ""


def _campos(hoja: _Hoja, filas: list[tuple[str, str, Any]], bloque: str, traza: dict) -> None:
    """Bloque de pares campo/valor con fuente y fecha de consulta de cada dato."""
    datos, rellenos = [], []
    for campo, etiqueta, valor in filas:
        fuente, fecha = _fuente(traza, bloque, campo)
        datos.append((etiqueta, _o_pendiente(valor), fuente, fecha))
        rellenos.append(_RELLENO_PENDIENTE if _vacio(valor) else None)
    hoja.tabla(["Campo", "Valor", "Fuente", "Fecha de consulta"], datos, rellenos=rellenos)


def _numero(valor: Any) -> float | None:
    """Monto registrado como número; vacío o ilegible -> None (celda vacía)."""
    try:
        return None if _vacio(valor) else float(valor)
    except (TypeError, ValueError):
        return None


def _porcentaje(valor: Any) -> float | None:
    """Participación registrada como 40 (%) -> 0.40 para el formato de Excel."""
    numero = _numero(valor)
    return None if numero is None else numero / 100


# ---------------------------------------------------------------------------
# Hojas
# ---------------------------------------------------------------------------

def _hoja_resumen(libro, audit, anio, validacion, cierre, risks, hojas: list[str], generado: str) -> None:
    hoja = _Hoja(libro, "Resumen", (34, 90))
    hoja.titulo("Levantamiento de información general del cliente", "Atlas · Auddit — resumen preliminar")
    hoja.tabla(["Dato", "Valor"], [
        ("Empresa", audit["company_name"]),
        ("RUC", _o_pendiente(audit["ruc"])),
        ("Período auditado", audit["period"]),
        ("Año fiscal de los EEFF", anio or PENDIENTE),
        ("Auditor asignado", _o_pendiente(row_get(audit, "auditor_name"))),
        ("Generado", generado),
    ])
    hoja.seccion("Estado del levantamiento")
    hoja.tabla(["Indicador", "Valor"], [
        ("Requisitos obligatorios cumplidos",
         f"{validacion['requisitos_cumplidos']} de {validacion['requisitos_total']}"),
        ("Pendientes obligatorios", len(cierre["pendientes"])),
        ("Recomendaciones pendientes", len(cierre["recomendados"])),
        ("Alertas y riesgos preliminares", len(risks)),
        ("Recomendación preliminar", cierre["recomendacion"]),
    ])
    hoja.seccion("Contenido del libro")
    hoja.tabla(["Hoja", "Contenido"], [(nombre, descripcion) for nombre, descripcion in hojas])
    _poner(hoja.ws, hoja.fila, 1, AVISO).font = Font(bold=True, color="9C5700")
    hoja.ws.merge_cells(start_row=hoja.fila, start_column=1, end_row=hoja.fila, end_column=2)
    hoja.ws.row_dimensions[hoja.fila].height = 32


def _hoja_personas(libro, titulo, encabezados, filas, formatos=None, total=None) -> None:
    hoja = _Hoja(libro, titulo, (18, 16, 38, 26, 18, 16, 26, 38, 16)[:len(encabezados)])
    hoja.titulo(titulo)
    primera, ultima = hoja.tabla(encabezados, filas, formatos, fijar_encabezado=True)
    if total and filas:
        fila = ultima + 1
        _poner(hoja.ws, fila, 1, "Total").font = Font(bold=True)
        for col, formato in total.items():
            letra = get_column_letter(col)
            cell = hoja.ws.cell(row=fila, column=col, value=f"=SUM({letra}{primera}:{letra}{ultima})")
            cell.number_format, cell.font = formato, Font(bold=True)
        for col in range(1, len(encabezados) + 1):
            hoja.ws.cell(row=fila, column=col).fill = _RELLENO_TOTAL


def _hoja_financiero(libro, snapshot, indicators, years, traza) -> None:
    hoja = _Hoja(libro, "6. Financiero", (52, 20, 20, 20, 20, 20, 20))
    anio = row_get(snapshot, "anio_fiscal", None)
    hoja.titulo("6. Información financiera", f"Estados financieros del año fiscal {anio}" if anio else "")
    fuente, fecha = _fuente(traza, BLOQUE_FINANCIERO, f"{anio}.activo_total") if anio else ("", "")
    hoja.tabla(["Dato", "Valor"], [
        ("Año fiscal", anio or PENDIENTE),
        ("Fecha de corte", _o_pendiente(row_get(snapshot, "fecha_corte"))),
        ("Fecha de aprobación de la junta", _o_pendiente(row_get(snapshot, "fecha_junta_aprobacion"))),
        ("Fuente", _o_pendiente(row_get(snapshot, "fuente") or fuente)),
        ("Fecha de consulta", _o_pendiente(row_get(snapshot, "fecha_consulta") or fecha)),
    ])
    hoja.seccion("Casilleros del estado financiero")
    if snapshot:
        filas = [
            (etiqueta + (" — calculado" if calculado else ""), valor if valor is not None else PENDIENTE)
            for etiqueta, valor, calculado in filas_comparativo(snapshot)
        ]
    else:
        filas = []
    hoja.tabla(["Concepto", "Valor (USD)"], filas, {2: _FORMATO_USD}, vacia="Sin datos financieros registrados.")
    hoja.seccion("Indicadores calculados")
    hoja.tabla(["Indicador", "Valor"], [
        ("Razón de endeudamiento (Pasivo / Activo)", indicators.get("razon_endeudamiento")),
        ("Margen neto (Utilidad / Ingresos)", indicators.get("margen_neto")),
        ("Patrimonio / Activo", indicators.get("patrimonio_sobre_activo")),
    ] if indicators.get("tiene_datos") else [], {2: _FORMATO_PCT}, vacia="Sin datos para calcular indicadores.")
    if len(years) > 1:
        hoja.seccion("Ejercicios registrados del RUC")
        anios = [row_get(y, "anio_fiscal") for y in years]
        por_anio = [filas_comparativo(y) for y in years]
        filas = [
            [etiqueta] + [conceptos[i][1] for conceptos in por_anio]
            for i, (etiqueta, _v, _c) in enumerate(por_anio[0])
        ]
        hoja.tabla(["Concepto", *(str(a) for a in anios)], filas,
                   {col: _FORMATO_USD for col in range(2, len(anios) + 2)})


def _hoja_requisitos(libro, validacion, cierre) -> None:
    hoja = _Hoja(libro, "Requisitos", (18, 70, 16, 16))
    hoja.titulo("Requisitos del levantamiento",
                "Obligatorio: si falta, queda pendiente. Recomendado: se sugiere completar.")
    filas, rellenos = [], []
    for req in validacion["requisitos"]:
        filas.append((req["source"], req["label"], "Obligatorio", "Cumplido" if req["ok"] else "Pendiente"))
        rellenos.append(None if req["ok"] else _RELLENO_PENDIENTE)
    requisitos = {r["label"] for r in validacion["requisitos"]}
    for pendiente in cierre["pendientes"]:
        if pendiente not in requisitos:  # tratamiento de alertas críticas, año fiscal
            filas.append(("Validaciones", pendiente, "Obligatorio", "Pendiente"))
            rellenos.append(_RELLENO_PENDIENTE)
    for rec in validacion["recomendaciones"]:
        filas.append((rec["source"], rec["label"], "Recomendado", "Pendiente"))
        rellenos.append(None)
    hoja.tabla(["Bloque", "Requisito", "Tipo", "Estado"], filas, rellenos=rellenos, fijar_encabezado=True)


def _hoja_validaciones(libro, validacion, risks) -> None:
    hoja = _Hoja(libro, "Validaciones y alertas", (42, 18, 14, 80))
    hoja.titulo("Validaciones cruzadas y alertas")
    hoja.seccion("Validaciones cruzadas")
    cruces = validacion["cruces"]
    hoja.tabla(
        ["Regla", "Resultado", "Nivel", "Detalle"],
        [(c["regla"], _ESTADOS_CRUCE.get(c["estado"], c["estado"]), _NIVELES.get(c["nivel"], c["nivel"]),
          c["detalle"]) for c in cruces],
        rellenos=[_RELLENO_ALERTA if c["estado"] == "no_coincide" else None for c in cruces],
    )
    hoja.seccion("Alertas automáticas")
    hoja.tabla(
        ["Nivel", "Alerta", "Tratamiento del auditor"],
        [(_NIVELES.get(a["nivel"], a["nivel"]), a["mensaje"], a["tratamiento"] or "Sin tratamiento registrado")
         for a in validacion["alertas"]],
        vacia="Sin alertas automáticas.",
    )
    hoja.seccion("Riesgos preliminares")
    hoja.tabla(["Riesgo"], [(r,) for r in risks], vacia="Sin riesgos preliminares con la información registrada.")


def _hoja_fuentes(libro, source_checks, sources, provenance) -> None:
    hoja = _Hoja(libro, "Fuentes y evidencia", (30, 44, 28, 40, 30, 16, 20))
    hoja.titulo("Fuentes consultadas y evidencia")
    hoja.seccion("Fuentes del levantamiento")
    hoja.tabla(
        ["Fuente", "Uso", "Estado", "Observación", "Consultada el"],
        [(row_get(s, "fuente"), row_get(s, "uso"),
          "Consultada" if row_get(s, "estado") == "consultada" else "Pendiente",
          row_get(s, "observacion") or "", row_get(s, "consultada_at") or "") for s in source_checks],
    )
    hoja.seccion("Evidencia registrada")
    hoja.tabla(
        ["Tipo", "Título", "URL", "Notas", "Registrada el"],
        [(row_get(s, "source_type") or "Fuente", row_get(s, "title") or "Sin título", row_get(s, "url") or "",
          row_get(s, "notes") or "", row_get(s, "created_at") or "") for s in sources],
    )
    hoja.seccion("Trazabilidad por dato")
    hoja.tabla(
        ["Bloque", "Campo", "Valor anterior", "Valor nuevo", "Fuente", "Fecha de consulta", "Registrado el"],
        [(_BLOQUES.get(row_get(p, "bloque"), row_get(p, "bloque")), row_get(p, "campo"),
          row_get(p, "valor_anterior") or "", row_get(p, "valor_nuevo") or "", row_get(p, "fuente"),
          row_get(p, "fecha_consulta"), row_get(p, "registrado_at")) for p in provenance],
    )


def _hoja_hallazgos(libro, data: dict) -> None:
    hoja = _Hoja(libro, "Hallazgos", (34, 100))
    hoja.titulo("Hallazgos preliminares e información complementaria")
    signals = extract_signals(data.get("pasted_text") or "")
    filas = [
        ("Observaciones del auditor", data.get("observations")),
        ("Riesgos ingresados manualmente", data.get("risk_flags")),
        ("Obligaciones tributarias", data.get("tax_obligations")),
        ("Contratación pública", data.get("public_contracting") or data.get("sercop_info")),
        ("RUC detectados en texto", ", ".join(signals["rucs"])),
        ("Correos detectados", ", ".join(signals["emails"])),
        ("Teléfonos detectados", ", ".join(signals["phones"])),
    ]
    hoja.tabla(["Hallazgo", "Detalle"], [(k, v) for k, v in filas if not _vacio(v)],
               vacia="Sin hallazgos adicionales registrados.")


# ---------------------------------------------------------------------------
# Libro completo
# ---------------------------------------------------------------------------

_HOJAS = [
    ("1. SRI", "Identificación tributaria"),
    ("2. Supercias", "Información societaria"),
    ("3. Ubicación", "Dirección de la compañía"),
    ("4. Administradores", "Nómina de administradores"),
    ("5. Accionistas", "Nómina de accionistas, capital y participación"),
    ("6. Financiero", "Estados financieros e indicadores"),
    ("Requisitos", "Cumplimiento de cada requisito del levantamiento"),
    ("Validaciones y alertas", "Cruces entre fuentes, alertas y riesgos"),
    ("Fuentes y evidencia", "Fuentes consultadas, evidencia y trazabilidad por dato"),
    ("Hallazgos", "Observaciones e información complementaria"),
]


def build_resumen_xlsx(audit: Any, ctx: dict[str, Any], generado: datetime | None = None) -> bytes:
    """Libro .xlsx del levantamiento de información del expediente.

    audit es la fila de get_audit() y ctx el contexto de get_audit_context().
    """
    research = ctx.get("research")
    data = {k: (research[k] or "") for k in research.keys()} if research is not None else {}
    profile = _CamposTolerantes(dict(ctx["profile"])) if ctx.get("profile") is not None else _CamposTolerantes()
    location = _CamposTolerantes(dict(ctx["location"])) if ctx.get("location") is not None else _CamposTolerantes()
    admins, shareholders = ctx.get("admins") or [], ctx.get("shareholders") or []
    snapshot = dict(ctx["snapshot"]) if ctx.get("snapshot") else None
    indicators = compute_indicators(snapshot) if snapshot else {}
    sources = ctx.get("sources") or []
    provenance = ctx.get("provenance") or []
    traza = _trazabilidad(provenance)

    validacion = evaluar_levantamiento(
        audit, profile, location, admins, shareholders, snapshot, ctx.get("alert_treatments"),
    )
    risks = risk_suggestions(audit, data, len(sources), profile, indicators, validacion)
    cierre = cierre_levantamiento(validacion, snapshot, risks)
    anio = row_get(snapshot, "anio_fiscal", None)

    libro = Workbook()
    libro.remove(libro.active)
    _hoja_resumen(libro, audit, anio, validacion, cierre, risks, _HOJAS,
                  (generado or datetime.now()).strftime("%Y-%m-%d %H:%M"))

    hoja = _Hoja(libro, "1. SRI", (30, 60, 34, 18))
    hoja.titulo("1. Identificación tributaria (SRI)")
    _campos(hoja, [("ruc", "RUC", audit["ruc"])] + [(c, l, profile[c]) for c, l in CAMPOS_SRI], BLOQUE_SRI, traza)

    hoja = _Hoja(libro, "2. Supercias", (30, 60, 34, 18))
    hoja.titulo("2. Información societaria (Supercias)")
    _campos(hoja, valores_supercias(profile, data) + [("ciiu_nivel6", "Código CIIU (Supercias)",
                                                      profile["ciiu_nivel6"])], BLOQUE_SUPERCIAS, traza)

    hoja = _Hoja(libro, "3. Ubicación", (30, 60, 34, 18))
    hoja.titulo("3. Ubicación")
    _campos(hoja, [(c, l, location[c]) for c, l in CAMPOS_UBICACION], BLOQUE_UBICACION, traza)

    _hoja_personas(
        libro, "4. Administradores",
        ["Identificación", "Tipo de identificación", "Nombre", "Cargo", "Nacionalidad", "Fuente",
         "Fecha de consulta"],
        [(_o_pendiente(row_get(a, "identificacion")), row_get(a, "tipo_identificacion") or "",
          row_get(a, "nombre"), _o_pendiente(row_get(a, "cargo")), _o_pendiente(row_get(a, "nacionalidad")),
          row_get(a, "fuente") or "", row_get(a, "fecha_consulta") or "") for a in admins],
    )
    _hoja_personas(
        libro, "5. Accionistas",
        ["N.º", "Identificación", "Nombre", "Tipo de identificación", "Capital (USD)", "Participación",
         "Beneficiario final", "Fuente", "Fecha de consulta"],
        [(row_get(s, "numero") or "", _o_pendiente(row_get(s, "identificacion")), row_get(s, "nombre"),
          row_get(s, "tipo_identificacion") or "", _numero(row_get(s, "capital")),
          _porcentaje(row_get(s, "participacion_porcentaje")), row_get(s, "beneficiario_final") or "",
          row_get(s, "fuente") or "", row_get(s, "fecha_consulta") or "") for s in shareholders],
        formatos={5: _FORMATO_USD, 6: _FORMATO_PCT},
        total={5: _FORMATO_USD, 6: _FORMATO_PCT},
    )
    _hoja_financiero(libro, snapshot, indicators, (ctx.get("financial") or {}).get("years") or [], traza)
    _hoja_requisitos(libro, validacion, cierre)
    _hoja_validaciones(libro, validacion, risks)
    _hoja_fuentes(libro, ctx.get("source_checks") or [], sources, provenance)
    _hoja_hallazgos(libro, data)

    libro.properties.title = f"Levantamiento de información — {audit['company_name']}"
    libro.properties.creator = "Atlas · Auddit"
    salida = io.BytesIO()
    libro.save(salida)
    return salida.getvalue()
