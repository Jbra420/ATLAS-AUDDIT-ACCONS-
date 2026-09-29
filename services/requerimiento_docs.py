"""
services/requerimiento_docs.py — Documentos del Requerimiento inicial.

Genera la carta de encargo y los dos certificados (PDF) y la solicitud inicial
de información (Excel de cuatro hojas), y lee las respuestas del Excel que
devuelve el cliente. Recibe los datos ya confirmados: no consulta la base ni
completa lo que falta.

Los textos se transcribieron de los ejemplos del documento de proceso de
Auddit (ver services.requerimiento.AVISO_PLANTILLA). Las firmas no se
reproducen: cada documento deja la línea para firmar.
"""
from __future__ import annotations

import io
import json
import re
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.packaging.custom import StringProperty
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

from services.normalizacion import normalizar_texto
from services.pdf_simple import PdfDocumento
from services.requerimiento import (
    CUADROS,
    FILAS_MINIMAS_CUADRO,
    HOJAS,
    ITEMS,
    ITEMS_HOJA1,
    ITEMS_HOJA2,
    fecha_larga,
    normalizar_nombre,
    representante_completo,
    texto_item,
)
from services.resumen_excel import _poner

LOGO_AUDDIT = Path(__file__).resolve().parent.parent / "static" / "logoauddit.jpeg"


def _anio_fin(datos: dict[str, Any]) -> str:
    return f"31 de diciembre del {datos['anio_auditado']}"


# ── Carta de encargo ─────────────────────────────────────────────────────

def build_carta(datos: dict[str, Any], logo: bytes | None = None) -> bytes:
    empresa = normalizar_nombre(datos["empresa"])
    representante = representante_completo(datos)
    doc = PdfDocumento()
    logo = logo if logo is not None else (LOGO_AUDDIT.read_bytes() if LOGO_AUDDIT.exists() else None)
    if logo:
        doc.imagen_jpeg(logo, 150, "derecha", despues=10)
    doc.parrafo(
        f"CARTA DE ENCARGO DE AUDITORÍA A LOS ESTADOS FINANCIEROS DE LA EMPRESA {empresa}, DEL AÑO QUE "
        f"TERMINA AL 31 DE DICIEMBRE DE {datos['anio_auditado']}",
        negrita=True, alineacion="centro", tamano=11, despues=18,
    )
    doc.parrafo(fecha_larga(datos["fecha_documentos"]), negrita=True, alineacion="derecha", despues=14)
    for linea in (f"A LA DIRECCION DE {empresa}", f"ATT. {representante}", "Representante Legal", "Ciudad."):
        doc.parrafo(linea, negrita=True, alineacion="izquierda", despues=0)
    doc.espacio(12)
    doc.parrafo("De nuestra consideración,", alineacion="izquierda")
    doc.parrafo(
        "Por medio de la presente, dejamos constancia de los términos y condiciones bajo los cuales se llevará a "
        f"cabo la auditoría a los Estados Financieros de {empresa}, correspondientes al ejercicio económico "
        f"terminado el {_anio_fin(datos)}."
    )

    def seccion(titulo: str, *parrafos: str) -> None:
        doc.espacio(4)
        doc.parrafo(titulo, negrita=True, alineacion="izquierda", despues=6)
        for texto in parrafos:
            doc.parrafo(texto)

    seccion(
        "1. Objetivo y alcance de la auditoría",
        f"El objetivo de nuestro encargo es realizar una auditoría de los estados financieros de {empresa} "
        f"correspondientes al ejercicio económico terminado el {_anio_fin(datos)} y expresar una opinión "
        "independiente sobre si éstos han sido preparados, en todos los aspectos materiales, de conformidad con "
        "el marco de información financiera aplicable.",
        "La auditoría comprenderá el examen del Estado de Situación Financiera, Estado del Resultado Integral, "
        "Estado de Cambios en el Patrimonio, Estado de Flujos de Efectivo y las correspondientes notas "
        "explicativas a los estados financieros.",
        "Nuestro trabajo será realizado de conformidad con las Normas Internacionales de Auditoría (NIA) y demás "
        "disposiciones legales y regulatorias aplicables en Ecuador.",
    )
    seccion(
        "2. Responsabilidad del auditor",
        "Nuestra responsabilidad consiste en expresar una opinión sobre los Estados Financieros con base en la "
        "auditoría realizada. Nuestro trabajo será planificado y ejecutado de conformidad con las Normas "
        "Internacionales de Auditoría (NIA), las cuales requieren el cumplimiento de lineamientos éticos y la "
        "obtención de seguridad razonable de que los Estados Financieros, considerados en su conjunto, estén "
        "libres de incorrecciones materiales debidas a fraude o error.",
        "Debido a las limitaciones inherentes de una auditoría y del control interno, existe un riesgo inevitable "
        "de que algunas incorrecciones materiales puedan no ser detectadas, aun cuando la auditoría haya sido "
        "adecuadamente planificada y ejecutada de conformidad con las NIA.",
    )
    seccion(
        "3. Responsabilidades de la administración",
        "La Administración reconoce y acepta su responsabilidad respecto de:",
        "a) La preparación y presentación razonable de los Estados Financieros de conformidad con el marco de "
        "información financiera aplicable.\n"
        "b) El diseño, implementación y mantenimiento del control interno que considere necesario para permitir "
        "la preparación de Estados Financieros libres de incorrecciones materiales, debidas a fraude o error.\n"
        "c) Proporcionar al equipo de auditoría, de manera completa y dentro de los plazos establecidos, acceso a "
        "toda la información relevante para la preparación de los Estados Financieros, así como cualquier "
        "información adicional que sea requerida para el desarrollo de la auditoría.\n"
        "d) Informar de manera oportuna sobre hechos, operaciones o circunstancias que pudieran afectar "
        "significativamente a los estados financieros.\n"
        "e) Comunicar al equipo de auditoría los hechos posteriores ocurridos entre la fecha de los Estados "
        "Financieros y la fecha del informe del auditor que pudieran requerir ajustes o revelación.\n"
        "f) Proporcionar manifestaciones escritas que sean requeridas como parte de la evidencia de auditoría "
        "obtenida durante el desarrollo del encargo.",
    )
    seccion("4. Equipo de auditoría",
            "El equipo de trabajo designado para la ejecución del encargo estará conformado por:")
    for integrante in datos["equipo"]:
        doc.vineta(integrante)
    seccion(
        "5. Planificación y cronograma",
        "El proceso de auditoria preliminar se realizará con información financiera con corte al "
        f"{fecha_larga(datos['fecha_corte'])}. El proceso relacionado "
        "con el levantamiento de inventarios o activos fijos se aplicará en fechas tentativas "
        f"{datos['fechas_inventario']}.",
        "El cronograma previsto de entrega de informes es el siguiente:",
    )
    doc.tabla(("Entregable", "Fecha de Entrega"), [(c["entregable"], c["fecha"]) for c in datos["cronograma"]])
    doc.parrafo(
        "Las fechas señaladas son estimadas y se encuentran sujetas al cierre definitivo de los Estados "
        "Financieros y a la entrega completa y oportuna de la información requerida por parte de la Administración."
    )
    doc.parrafo(
        "Durante el proceso de auditoría se comunicarán oportunamente a la Administración las situaciones "
        "relevantes identificadas que, de acuerdo con nuestro juicio profesional, requieran su conocimiento y "
        "atención"
    )
    seccion(
        "6. Informes a emitir",
        "Como resultado del encargo, se prevé la emisión del correspondiente Informe del Auditor Independiente "
        "sobre los Estados Financieros y demás informes establecidos en el alcance contratado y en la normativa "
        "aplicable.",
        "La estructura y contenido del informe del auditor se prepararán de conformidad con las Normas "
        "Internacionales de Auditoría. No obstante, dependiendo de los resultados y circunstancias identificadas "
        "durante el desarrollo del trabajo, la forma y contenido del informe finalmente emitido podrían diferir de "
        "los inicialmente previstos.",
    )
    seccion(
        "7. Aceptación de los términos del encargo",
        "Solicitamos se sirva manifestar su aceptación de los términos establecidos en la presente carta de "
        "encargo mediante la firma del presente documento.",
        "La suscripción de este documento evidencia el reconocimiento y aceptación de las responsabilidades de "
        "la Administración y del auditor descritas anteriormente.",
    )
    doc.firmas([
        [representante, normalizar_nombre(datos["representante_cargo"]), empresa],
        [normalizar_nombre(datos["auddit_representante"]), normalizar_nombre(datos["auddit_cargo"]),
         "AUDDIT S.A.S."],
    ])
    return doc.bytes()


# ── Certificados ─────────────────────────────────────────────────────────

_OBJETO_CERTIFICADO = {
    "cert_relacionadas": "con empresas relacionadas locales y/o empresas relacionadas del exterior",
    "cert_paraisos": "con empresas relacionadas domiciliadas en jurisdicciones consideradas Paraísos Fiscales.",
}


def build_certificado(tipo: str, datos: dict[str, Any]) -> bytes:
    """Certificado que firma el representante del cliente (sin marcar SI/NO:
    lo responde el cliente)."""
    empresa = normalizar_nombre(datos["empresa"])
    representante = representante_completo(datos)
    identificacion = datos["representante_identificacion"]
    documento = "cédula de identidad" if identificacion.isdigit() and len(identificacion) == 10 else "pasaporte"
    doc = PdfDocumento(margen_x=72)
    doc.parrafo("CERTIFICADO", negrita=True, alineacion="centro", tamano=11, despues=26)
    doc.parrafo(fecha_larga(datos["fecha_documentos"]), negrita=True, alineacion="derecha", despues=26)
    for linea in ("A LA DIRECCION DE AUDDIT S.A.S.", f"ATT. {normalizar_nombre(datos['auddit_representante'])}",
                  "Representante Legal", "Ciudad."):
        doc.parrafo(linea, negrita=True, alineacion="izquierda", despues=0)
    doc.espacio(26)
    doc.parrafo("De mi consideración:", negrita=True, alineacion="izquierda", despues=20)
    doc.parrafo(
        f"Yo, {representante} con {documento} {identificacion}, de nacionalidad "
        f"{datos['representante_nacionalidad'].lower()}, mayor de edad, domiciliado(a) en la ciudad de "
        f"{datos['representante_ciudad']}, en mi calidad de {normalizar_nombre(datos['representante_cargo'])} de "
        f"{empresa}, con RUC {datos['ruc']}, comparezco ante ustedes con el fin de CERTIFICAR que en el año "
        f"{datos['anio_certificados']} la empresa que represento, SI( ) NO( ) mantiene relaciones comerciales, "
        f"financieras o contractuales {_OBJETO_CERTIFICADO[tipo]}",
        interlineado=1.6,
    )
    doc.parrafo("Particular que pongo en su conocimiento para los fines pertinentes.", alineacion="izquierda",
                despues=34)
    doc.parrafo("Atentamente:", negrita=True, alineacion="izquierda")
    doc.firmas([[representante, normalizar_nombre(datos["representante_cargo"]), empresa]], espacio_firma=70)
    return doc.bytes()


# ── Solicitud inicial de información (Excel) ────────────────────────────

_AZUL = "1F3864"
_MAGENTA = "C000C0"
_ENCABEZADO = Font(bold=True, color="FFFFFF")
_RELLENO_AZUL = PatternFill("solid", fgColor=_AZUL)
_RELLENO_MAGENTA = PatternFill("solid", fgColor=_MAGENTA)
_LINEA = Side(style="thin", color="7F7F7F")
_BORDE = Border(left=_LINEA, right=_LINEA, top=_LINEA, bottom=_LINEA)
_CENTRO = Alignment(horizontal="center", vertical="center", wrap_text=True)
_INSTRUCCION = ("Utilice el presente formato para especificar la información que la empresa no aplica y "
                "controlar la información que ya esta lista para ser enviado al equipo de auditoria AUDDIT.")
_CIERRE = ("De no contar o tener alguna inquietud a la información antes solicitada, favor nos informen a fin de "
           "coordinar o tomar alguna decisión al respecto.")
MARCA = "X"


def _celda(ws, fila: int, col: int, valor: Any, *, negrita: bool = False, borde: bool = True,
           centro: bool = False, relleno: PatternFill | None = None, fuente: Font | None = None):
    cell = _poner(ws, fila, col, valor)
    if fuente is not None:
        cell.font = fuente
    elif negrita:
        cell.font = Font(bold=True)
    if borde:
        cell.border = _BORDE
    if centro:
        cell.alignment = _CENTRO
    if relleno is not None:
        cell.fill = relleno
    return cell


def _titulo(ws, fila: int, texto: str, ultima: int, *, tamano: int = 12, relleno=None, fuente=None) -> None:
    ws.merge_cells(start_row=fila, start_column=1, end_row=fila, end_column=ultima)
    _celda(ws, fila, 1, texto, borde=relleno is not None, centro=True, relleno=relleno,
           fuente=fuente or Font(bold=True, size=tamano))


# Identificación del libro: cada Excel generado lleva la referencia de su
# generación en las propiedades del documento y en una celda visible de las
# hojas 1 y 2 (sin agregar hojas). Al importar la respuesta se contrastan con
# lo registrado en Atlas; el nombre del archivo no cuenta.
PROPIEDADES_REF = {"AtlasReferencia": "referencia", "AtlasAuditoria": "audit_id", "AtlasRUC": "ruc",
                   "AtlasEjercicio": "ejercicio"}
_ETIQUETA_REF = "Referencia Atlas:"
_REF_RE = re.compile(r"Referencia Atlas:\s*(REQ-\d+-\d+-[0-9A-F]+)")
_SIN_REF = "VISTA PREVIA — sin referencia de Atlas: este archivo no se puede enviar ni importar"


def encabezado_empresa(datos: dict[str, Any]) -> str:
    return f"{normalizar_nombre(datos['empresa'])} — RUC {datos['ruc']} — Ejercicio {datos['anio_auditado']}"


def _cabecera_hoja(ws, numero: int, datos: dict[str, Any], seccion: str, ultima: int, referencia: str) -> int:
    _titulo(ws, 1, "AUDDIT S.A.S.", ultima, tamano=16, fuente=Font(bold=True, size=16, color="5B2A86"))
    _titulo(ws, 2, f"{_ETIQUETA_REF} {referencia} (no modifique esta celda)" if referencia else _SIN_REF, ultima,
            tamano=8, fuente=Font(size=8, color="7F7F7F"))
    _titulo(ws, 3, f"REQUERIMIENTO DE INFORMACIÓN #{numero}", ultima)
    _titulo(ws, 4, encabezado_empresa(datos), ultima, tamano=10, fuente=Font(size=10))
    ws.merge_cells(start_row=5, start_column=1, end_row=5, end_column=ultima)
    _celda(ws, 5, 1, _INSTRUCCION, borde=False, centro=True)
    ws.row_dimensions[5].height = 32
    ws.cell(row=7, column=1, value=seccion).font = Font(bold=True)
    return 8


def _validacion_marca(ws, rango: str) -> None:
    dv = DataValidation(type="list", formula1=f'"{MARCA}"', allow_blank=True,
                        error="Escriba X para marcar", errorTitle="Marca no válida")
    dv.add(rango)
    ws.add_data_validation(dv)


def _hoja_items(wb: Workbook, hoja: int, datos, marcas, anclas: dict[str, tuple[int, int]], referencia: str) -> None:
    ws = wb.active if hoja == 1 else wb.create_sheet()
    ws.title = HOJAS[hoja]
    con_formato = hoja == 2
    columnas = ("Nro.", "Detalle", *(("Formato",) if con_formato else ()), "CUMPLIDO", "NO APLICA", "OBSERVACIONES")
    anchos = (7, 70, *((10,) if con_formato else ()), 12, 12, 30)
    for i, ancho in enumerate(anchos, start=1):
        ws.column_dimensions[chr(64 + i)].width = ancho
    seccion = "1.- DOCUMENTOS DE LA COMPAÑÍA" if hoja == 1 else "2.- INFORMACION CONTABLE - TRIBUTARIA"
    fila = _cabecera_hoja(ws, hoja, datos, seccion, len(columnas), referencia)
    for col, texto in enumerate(columnas, start=1):
        _celda(ws, fila, col, texto, centro=True, relleno=_RELLENO_AZUL, fuente=_ENCABEZADO)
    c_cumplido = columnas.index("CUMPLIDO") + 1
    primera = fila + 1

    def item(numero: int, texto: str, formato: str, enlace: str) -> None:
        nonlocal fila
        fila += 1
        marca = marcas.get((hoja, numero)) or {}
        _celda(ws, fila, 1, numero, centro=True)
        detalle = _celda(ws, fila, 2, texto_item(texto, datos))
        if enlace in anclas:
            nombre, fila_ancla = anclas[enlace]
            detalle.hyperlink = f"#'{nombre}'!A{fila_ancla}"
            detalle.font = Font(color="0563C1", underline="single")
        if con_formato:
            _celda(ws, fila, 3, formato, centro=True)
        _celda(ws, fila, c_cumplido, MARCA if marca.get("cumplido") else None, centro=True)
        _celda(ws, fila, c_cumplido + 1, MARCA if marca.get("no_aplica") else None, centro=True)
        _celda(ws, fila, c_cumplido + 2, marca.get("observacion") or None)
        ws.row_dimensions[fila].height = 32

    if hoja == 1:
        for numero, texto, enlace in ITEMS_HOJA1:
            item(numero, texto, "", enlace)
    else:
        for titulo, filas in ITEMS_HOJA2:
            fila += 1
            ws.merge_cells(start_row=fila, start_column=1, end_row=fila, end_column=len(columnas))
            _celda(ws, fila, 1, titulo, centro=True, relleno=_RELLENO_AZUL, fuente=_ENCABEZADO)
            for numero, texto, formato, enlace in filas:
                item(numero, texto, formato, enlace)
    letra = chr(64 + c_cumplido)
    _validacion_marca(ws, f"{letra}{primera}:{chr(64 + c_cumplido + 1)}{fila}")
    fila += 1
    ws.merge_cells(start_row=fila, start_column=1, end_row=fila, end_column=len(columnas))
    _celda(ws, fila, 1, _CIERRE, negrita=True, centro=True)
    ws.row_dimensions[fila].height = 32
    ws.freeze_panes = ws.cell(row=primera, column=1)


def _hoja_cuadros(wb: Workbook, hoja: int, detalles: dict[str, list[list[str]]]) -> dict[str, tuple[str, int]]:
    ws = wb.create_sheet(HOJAS[hoja])
    for i, ancho in enumerate((10, 34, 26, 20, 26, 28), start=1):
        ws.column_dimensions[chr(64 + i)].width = ancho
    anclas = {}
    fila = 1
    for clave, (hoja_cuadro, titulo, columnas, numerado) in CUADROS.items():
        if hoja_cuadro != hoja:
            continue
        cabecera = (("No.",) if numerado else ()) + columnas
        anclas[clave] = (HOJAS[hoja], fila)
        ws.merge_cells(start_row=fila, start_column=1, end_row=fila, end_column=len(cabecera))
        _celda(ws, fila, 1, titulo, centro=True, relleno=_RELLENO_MAGENTA, fuente=_ENCABEZADO)
        fila += 1
        for col, texto in enumerate(cabecera, start=1):
            _celda(ws, fila, col, texto.upper(), centro=True, relleno=_RELLENO_AZUL, fuente=_ENCABEZADO)
        filas = detalles.get(clave) or []
        for i in range(max(len(filas) + 2, FILAS_MINIMAS_CUADRO)):
            fila += 1
            valores = filas[i] if i < len(filas) else [None] * len(columnas)
            if numerado:
                _celda(ws, fila, 1, i + 1, centro=True)
            for col, valor in enumerate(valores, start=2 if numerado else 1):
                _celda(ws, fila, col, valor or None)
        fila += 3
    return anclas


def build_solicitud_xlsx(
    datos: dict[str, Any], marcas: dict[tuple[int, int], Any], detalles: dict[str, list[list[str]]],
    identificacion: dict[str, Any] | None = None,
) -> bytes:
    """Libro de cuatro hojas: requerimientos #1 y #2 con CUMPLIDO / NO APLICA /
    OBSERVACIONES, y los cuadros de detalle que completa el cliente.
    identificacion ({referencia, audit_id}) liga el libro a su generación; sin
    ella (vista previa) el libro dice que no se puede importar."""
    wb = Workbook()
    referencia = (identificacion or {}).get("referencia") or ""
    if referencia:
        valores = {"referencia": referencia, "audit_id": identificacion["audit_id"], "ruc": datos["ruc"],
                   "ejercicio": datos["anio_auditado"]}
        for nombre, clave in PROPIEDADES_REF.items():
            wb.custom_doc_props.append(StringProperty(name=nombre, value=str(valores[clave])))
    marcas = {k: {"cumplido": m["cumplido"], "no_aplica": m["no_aplica"], "observacion": m["observacion"]}
              for k, m in marcas.items()}
    # Las hojas de cuadros se escriben primero en memoria para conocer sus anclas.
    borrador = Workbook()
    anclas = {**_hoja_cuadros(borrador, 3, detalles), **_hoja_cuadros(borrador, 4, detalles)}
    _hoja_items(wb, 1, datos, marcas, anclas, referencia)
    _hoja_items(wb, 2, datos, marcas, anclas, referencia)
    _hoja_cuadros(wb, 3, detalles)
    _hoja_cuadros(wb, 4, detalles)
    salida = io.BytesIO()
    wb.save(salida)
    return salida.getvalue()


# ── Lectura del Excel respondido ────────────────────────────────────────

_MARCADO = {"X", "SI", "SÍ", "✓", "✔", "☑", "☒", "1", "TRUE", "VERDADERO"}
_SIN_MARCA = {"", "NO", "☐", "0", "FALSE", "FALSO"}


def _marca(valor: Any) -> bool | None:
    """True si la celda está marcada, False si está vacía; None si no se reconoce."""
    if valor is None or valor is False:
        return False
    if valor is True:
        return True
    texto = str(valor).strip().upper()
    if texto in _MARCADO:
        return True
    if texto in _SIN_MARCA:
        return False
    return None


def _texto_celda(valor: Any) -> str:
    """Valor como texto, sin evaluar nada: una fórmula se conserva como texto."""
    if valor is None:
        return ""
    if isinstance(valor, float) and valor.is_integer():
        valor = int(valor)
    return str(valor)


def leer_respuesta_xlsx(contenido: bytes) -> dict[str, Any]:
    """Lee marcas, observaciones y cuadros del Excel que devolvió el cliente.

    CUMPLIDO y NO APLICA son excluyentes; una marca no reconocida o un ítem
    marcado dos veces invalida la importación (el archivo recibido se conserva
    igual). Las observaciones se guardan tal como llegaron.
    """
    try:
        wb = load_workbook(io.BytesIO(contenido), read_only=True, data_only=False)
    except Exception as exc:  # noqa: BLE001 — archivo del cliente, cualquier formato inválido
        raise ValueError("No se pudo leer el Excel: el archivo está dañado o no es .xlsx") from exc
    errores: list[str] = []
    items: list[dict[str, Any]] = []
    if wb.sheetnames != [HOJAS[h] for h in sorted(HOJAS)]:
        errores.append("El libro debe tener exactamente las cuatro hojas del requerimiento, en su orden: "
                       + ", ".join(f"«{HOJAS[h]}»" for h in sorted(HOJAS)))
    identificacion: dict[str, Any] = {
        "propiedades": {p.name: str(p.value).strip() for p in wb.custom_doc_props if p.name in PROPIEDADES_REF},
        "celdas": {}, "encabezados": {}, "textos": {},
    }
    for hoja in (1, 2):
        if HOJAS[hoja] not in wb.sheetnames:
            errores.append(f"Falta la hoja «{HOJAS[hoja]}»")
            continue
        filas = list(wb[HOJAS[hoja]].iter_rows(values_only=True))
        primera = [_texto_celda(f[0]).strip() if f else "" for f in filas[:4]] + [""] * 4
        identificacion["celdas"][hoja] = primera[1]
        identificacion["encabezados"][hoja] = primera[3]
        cabecera = next((i for i, f in enumerate(filas) if any(_texto_celda(v).strip().upper() == "CUMPLIDO"
                                                                for v in f)), None)
        if cabecera is None:
            errores.append(f"Hoja «{HOJAS[hoja]}»: no se encontró la columna CUMPLIDO")
            continue
        nombres = [_texto_celda(v).strip().upper() for v in filas[cabecera]]
        c_cum, c_na, c_obs = (nombres.index(n) for n in ("CUMPLIDO", "NO APLICA", "OBSERVACIONES"))
        validos = {n for n, _t in ITEMS[hoja]}
        vistos = set()
        for fila in filas[cabecera + 1:]:
            numero = fila[0] if fila else None
            if isinstance(numero, str) and numero.strip().isdigit():
                numero = int(numero.strip())
            if isinstance(numero, float) and numero.is_integer():
                numero = int(numero)
            if not isinstance(numero, int) or numero not in validos:
                continue
            vistos.add(numero)
            identificacion["textos"][(hoja, numero)] = _texto_celda(fila[1] if len(fila) > 1 else None).strip()
            celdas = list(fila) + [None] * (max(c_cum, c_na, c_obs) + 1 - len(fila))
            cumplido, no_aplica = _marca(celdas[c_cum]), _marca(celdas[c_na])
            if cumplido is None or no_aplica is None:
                errores.append(f"Req. #{hoja} ítem {numero}: marca no reconocida (use X)")
                continue
            if cumplido and no_aplica:
                errores.append(f"Req. #{hoja} ítem {numero}: CUMPLIDO y NO APLICA no pueden marcarse juntos")
                continue
            items.append({"hoja": hoja, "numero": numero, "cumplido": cumplido, "no_aplica": no_aplica,
                          "observacion": _texto_celda(celdas[c_obs])})
        faltan = sorted(validos - vistos)
        if faltan:
            errores.append(f"Req. #{hoja}: faltan los ítems {', '.join(map(str, faltan))}")
    detalles = _leer_cuadros(wb, errores)
    wb.close()
    if errores:
        raise ValueError("El Excel no se importó: " + "; ".join(errores[:10])
                         + ("…" if len(errores) > 10 else ""))
    return {"items": items, "detalles": detalles, "identificacion": identificacion}


# ── Procedencia del Excel respondido ────────────────────────────────────

def referencia_del_libro(identificacion: dict[str, Any]) -> tuple[set[str], str]:
    """Referencias que trae el libro (propiedad y celdas visibles) y un error
    si se contradicen. Vacío si el libro no trae ninguna."""
    refs = set()
    if identificacion["propiedades"].get("AtlasReferencia"):
        refs.add(identificacion["propiedades"]["AtlasReferencia"])
    for texto in identificacion["celdas"].values():
        encontrada = _REF_RE.search(texto or "")
        if encontrada:
            refs.add(encontrada.group(1))
        elif texto and _ETIQUETA_REF in texto:
            return refs, "la celda de referencia de Atlas fue modificada"
    if len(refs) > 1:
        return refs, "las referencias de Atlas del libro no coinciden entre sí"
    return refs, ""


def verificar_procedencia(
    identificacion: dict[str, Any], paquete: Any, *, audit_id: int, ruc_auditoria: str, enviados: set[int],
) -> tuple[str, str]:
    """(estado, motivo) de un Excel respondido ya leído frente al registro de
    Atlas. paquete es la generación cuya referencia trae el libro (None si no
    trae o no existe). Estados: "importado", "rechazado" (ajeno, alterado o
    incoherente) y "revision_manual" (plausible pero no verificable)."""
    refs, error = referencia_del_libro(identificacion)
    if error:
        return "rechazado", f"El libro fue alterado: {error}"
    ruc_auditoria = (ruc_auditoria or "").strip()
    if not refs:
        encabezado = " ".join(identificacion["encabezados"].values())
        if ruc_auditoria and f"RUC {ruc_auditoria}" not in encabezado:
            return "rechazado", ("El Excel no trae la referencia de Atlas y su encabezado no corresponde al RUC "
                                 f"{ruc_auditoria} de esta auditoría")
        return "revision_manual", ("El Excel no trae la referencia de Atlas (libro anterior a la identificación, "
                                   "vista previa o contenido copiado a otro archivo): no se puede saber a qué "
                                   "generación responde. Revíselo manualmente o pida al cliente responder sobre "
                                   "el Excel enviado")
    if paquete is None:
        return "rechazado", "La referencia del Excel no corresponde a ninguna generación registrada en Atlas"
    if int(paquete["audit_id"]) != int(audit_id):
        return "rechazado", "El Excel corresponde a otra auditoría (otra empresa o ejercicio)"
    datos = json.loads(paquete["datos_json"])
    esperado = {"referencia": paquete["referencia"], "audit_id": str(audit_id), "ruc": str(datos["ruc"]),
                "ejercicio": str(datos["anio_auditado"])}
    for nombre, clave in PROPIEDADES_REF.items():
        valor = identificacion["propiedades"].get(nombre)
        if valor is not None and valor != esperado[clave]:
            return "rechazado", f"El dato «{clave}» del libro ({valor}) no coincide con Atlas ({esperado[clave]})"
    if ruc_auditoria and datos["ruc"] != ruc_auditoria:
        return "rechazado", f"El RUC del Excel ({datos['ruc']}) no es el de esta auditoría ({ruc_auditoria})"
    encabezado = encabezado_empresa(datos)
    for hoja, texto in identificacion["encabezados"].items():
        if " ".join(texto.split()) != encabezado:
            return "rechazado", (f"El encabezado de la hoja «{HOJAS[hoja]}» ({texto or 'vacío'}) no coincide con "
                                 f"la generación {paquete['numero']}: {encabezado}")
    for hoja, filas in ITEMS.items():
        for numero, plantilla in filas:
            recibido = " ".join((identificacion["textos"].get((hoja, numero)) or "").split())
            if recibido != " ".join(texto_item(plantilla, datos).split()):
                return "rechazado", (f"Req. #{hoja} ítem {numero}: el detalle no coincide con la generación "
                                     f"{paquete['numero']} (libro alterado o de otra generación)")
    if int(paquete["id"]) not in enviados:
        return "revision_manual", (f"Corresponde a la generación {paquete['numero']}, que no consta como enviada "
                                   "al cliente. Registre el envío de esa generación y vuelva a cargar el archivo, "
                                   "o revíselo manualmente")
    return "importado", f"Corresponde a la generación {paquete['numero']} enviada al cliente"


def _leer_cuadros(wb, errores: list[str]) -> dict[str, list[list[str]]]:
    titulos = {normalizar_texto(t): clave for clave, (_h, t, _c, _n) in CUADROS.items()}
    detalles: dict[str, list[list[str]]] = {clave: [] for clave in CUADROS}
    for hoja in (3, 4):
        if HOJAS[hoja] not in wb.sheetnames:
            errores.append(f"Falta la hoja «{HOJAS[hoja]}»")
            continue
        actual, saltar_cabecera = None, False
        for fila in wb[HOJAS[hoja]].iter_rows(values_only=True):
            textos = [_texto_celda(v).strip() for v in fila]
            primero = next((t for t in textos if t), "")
            if normalizar_texto(primero) in titulos:
                actual, saltar_cabecera = titulos[normalizar_texto(primero)], True
                continue
            if actual is None:
                continue
            if saltar_cabecera:
                saltar_cabecera = False
                continue
            _h, _t, columnas, numerado = CUADROS[actual]
            valores = textos[1:len(columnas) + 1] if numerado else textos[:len(columnas)]
            if any(valores):
                detalles[actual].append(valores + [""] * (len(columnas) - len(valores)))
    return detalles
