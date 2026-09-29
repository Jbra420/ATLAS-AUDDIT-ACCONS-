"""Pestaña principal "Requerimiento inicial" (segundo paso del proceso).

Base y carpeta de adjuntos temporales: nada toca auddit.db ni adjuntos/.
Las acciones del router y la vista se enlazan a la base temporal con
mock.patch, igual que en tests/test_integridad_expediente.py.
"""
from __future__ import annotations

import functools
import hashlib
import io
import json
import tempfile
import unittest
from concurrent.futures import ThreadPoolExecutor
from datetime import date, datetime, timedelta
from pathlib import Path
from unittest import mock

import pypdf
from openpyxl import load_workbook

import database.expedientes as expedientes
import database.requerimiento as R
from core import router
from core.server import FormData, content_disposition
from database import add_administrator, authenticate, connect, create_company_audit, create_user, init_db
from services.requerimiento import (
    CUADROS,
    DOCUMENTOS,
    HOJAS,
    ITEMS,
    faltantes,
    normalizar_datos,
    precarga,
    validar_archivo,
)
from services.pdf_simple import PdfDocumento
from services.requerimiento_docs import build_solicitud_xlsx, leer_respuesta_xlsx
from views.auditor import requerimiento as vista
from views.auditor.radar import page as radar_page

CEDULA = "0104926555"


def _pdf_legible(texto: str = "Documento firmado de prueba") -> bytes:
    doc = PdfDocumento()
    doc.parrafo(texto)
    return doc.bytes()


PDF = _pdf_legible()
PNG = b"\x89PNG\r\n\x1a\n" + b"0" * 16

DATOS = {
    "empresa": "Constructora Esgingenieria S.A.S.", "ruc": "0190377210001", "representante_titulo": "Mgtr.",
    "representante_nombre": "Eduardo Alfonso Serpa Garcia", "representante_cargo": "Gerente",
    "representante_identificacion": CEDULA, "representante_nacionalidad": "Ecuatoriana",
    "representante_ciudad": "Cuenca", "anio_auditado": "2026", "anio_certificados": "2025", "anio_cerrado": "2025",
    "fecha_documentos": "2026-09-01", "fecha_corte": "2026-07-31",
    "fechas_inventario": "entre el 15 de octubre y el 15 de diciembre",
    "auddit_representante": "Mgtr. Fernando Parra Suarez", "auddit_cargo": "Gerente",
    "correo_para_nombre": "Contadora", "correo_para": "contadora@cliente.ec", "correo_cc": "fparra@accons.ec",
    "equipo": "Mgtr. Fernando Parra Suarez\nLcda. Camila Guevara Lucero",
    **{f"cronograma_{i}": f"Hasta {m} 2027" for i, m in enumerate(("febrero", "marzo", "abril", "julio"))},
}


def _form(campos: dict | None = None, archivos: dict | None = None, **extra) -> FormData:
    todos = {**(campos or {}), **extra}
    return FormData({k: [str(v)] for k, v in todos.items()}, archivos or {})


class _Caso(unittest.TestCase):
    def setUp(self):
        tmp = Path(tempfile.mkdtemp())
        self.db, self.adjuntos = tmp / "atlas.db", tmp / "adjuntos"
        init_db(self.db, demo=True)
        self.admin = authenticate("admin", "admin123", self.db)
        self.auditor = authenticate("auditor", "auditor123", self.db)
        with connect(self.db) as conn:
            create_user(conn, "otro.auditor", "Otro Auditor", "auditor", "otra-clave-1")
        self.otro = authenticate("otro.auditor", "otra-clave-1", self.db)
        self.audit_id = create_company_audit(
            "CONSTRUCTORA ESGINGENIERIA S.A.S.", "0190377210001", "Cuenca", "", "2026",
            self.auditor["id"], self.admin["id"], self.db,
        )
        db, adj = self.db, self.adjuntos
        enlaces = {
            "get_audit": lambda i, u: expedientes.get_audit(i, u, db),
            "get_requerimiento_context": lambda a: R.get_requerimiento_context(a, db),
            "get_requerimiento_file": lambda a, o, i: R.get_requerimiento_file(a, o, i, db),
            "save_requerimiento_adjunto": functools.partial(R.save_requerimiento_adjunto, db_path=db, adjuntos_dir=adj),
            "save_requerimiento_datos": lambda a, d, u: R.save_requerimiento_datos(a, d, u, db),
            "save_requerimiento_items": lambda a, m: R.save_requerimiento_items(a, m, db),
            "save_requerimiento_detalle": lambda a, s, f: R.save_requerimiento_detalle(a, s, f, db),
            "get_requerimiento_paquete": lambda r: R.get_requerimiento_paquete(r, db),
            "next_requerimiento_numero": lambda a: R.next_requerimiento_numero(a, db),
            "register_requerimiento_paquete": functools.partial(
                R.register_requerimiento_paquete, db_path=db, adjuntos_dir=adj),
            "review_requerimiento_adjunto": lambda a, i, r, n, u: R.review_requerimiento_adjunto(a, i, r, n, u, db),
            "register_requerimiento_envio": functools.partial(R.register_requerimiento_envio, db_path=db),
            "save_requerimiento_importacion": functools.partial(R.save_requerimiento_importacion, db_path=db),
            "ruta_archivo": lambda r: R.ruta_archivo(r, adj),
        }
        vistas = {
            "get_audit": enlaces["get_audit"],
            "get_requerimiento_context": enlaces["get_requerimiento_context"],
            "get_audit_context": lambda a: expedientes.get_audit_context(a, db),
        }
        for destino, cambios in ((router, enlaces), (vista, vistas)):
            parche = mock.patch.multiple(destino, **cambios)
            parche.start()
            self.addCleanup(parche.stop)

    # ── Atajos ───────────────────────────────────────────────────────────
    @property
    def audit(self):
        return expedientes.get_audit(self.audit_id, self.auditor, self.db)

    def accion(self, ruta: str, form: FormData) -> str:
        return router.REQUERIMIENTO_POSTS[ruta][1](form, self.audit, self.auditor)

    def contrato(self, revisar: bool = True):
        mensaje = self.accion("/auditor/requerimiento/contrato", _form(audit_id=self.audit_id,
                              archivos={"archivo": [("contrato.pdf", PDF)]}))
        if revisar:
            self.revisar(self.contexto()["adjuntos"]["contrato"][0]["id"])
        return mensaje

    def revisar(self, adjunto_id: int, resultado: str = "conforme", nota: str = "", todas: bool = True) -> str:
        checks = {f"verif_{c}": "1" for c in ("empresa", "ejercicio", "integridad", "firmas")} if todas else {}
        return self.accion("/auditor/requerimiento/revisar",
                           _form(adjunto_id=adjunto_id, resultado=resultado, nota=nota, **checks))

    def confirmar(self, **cambios):
        return self.accion("/auditor/requerimiento/datos", _form({**DATOS, **cambios}))

    def generar(self) -> str:
        return self.accion("/auditor/requerimiento/generar", _form())

    def contexto(self) -> dict:
        return R.get_requerimiento_context(self.audit_id, self.db)

    def pagina(self, user) -> str:
        return vista.render(user, {"audit_id": [str(self.audit_id)]}, "/auditor/requerimiento", csrf_token="t")

    def listo(self):
        self.contrato()
        self.confirmar()
        self.generar()


class TestNavegacion(_Caso):
    def test_url_de_retorno_codifica_el_fragmento(self):
        url = router.requerimiento_url(self.audit_id, "correo\r\nX-Test: si", err="Inválido")
        self.assertNotIn("\r", url)
        self.assertNotIn("\n", url)
        self.assertIn("#correo%0D%0AX-Test%3A+si", url)

    def test_levantamiento_y_requerimiento_son_pantallas_principales_separadas(self):
        with mock.patch.multiple(radar_page, get_audit=lambda i, u: expedientes.get_audit(i, u, self.db),
                                 get_audit_context=lambda a: expedientes.get_audit_context(a, self.db)):
            radar = radar_page.render(self.auditor, {"audit_id": [str(self.audit_id)]}, "/auditor/radar", "t")
        self.assertIn("<title>Levantamiento de información | Atlas</title>", radar)
        self.assertIn(f'href="/auditor/requerimiento?audit_id={self.audit_id}"', radar)
        self.assertNotIn('id="tab-requerimiento"', radar, "No es una novena pestaña del levantamiento")
        for pestana in ("sri", "supercias", "ubicacion", "admins", "accionistas", "indicadores", "documentos", "resumen"):
            self.assertIn(f'id="tab-{pestana}"', radar)

        pagina = self.pagina(self.auditor)
        self.assertIn("<title>Requerimiento inicial | Atlas</title>", pagina)
        self.assertIn(f'href="/auditor/radar?audit_id={self.audit_id}"', pagina)
        self.assertIn('class="proceso-tab active" href="/auditor/requerimiento', pagina)
        self.assertNotIn("radar-tab-pane", pagina)

    def test_el_jefe_navega_a_sus_rutas_de_solo_lectura(self):
        pagina = self.pagina(self.admin)
        self.assertIn(f'href="/admin/audit?audit_id={self.audit_id}"', pagina)
        self.assertIn(f'href="/admin/requerimiento?audit_id={self.audit_id}"', pagina)

    def test_el_estado_se_muestra_sin_contrato(self):
        pagina = self.pagina(self.auditor)
        for paso in ("Contrato firmado y revisado", "Datos del requerimiento confirmados",
                     "Documentos generados con los datos vigentes", "Correo enviado al cliente",
                     "Cliente notificado por WhatsApp", "Documentos del cliente recibidos y revisados"):
            self.assertIn(paso, pagina)
        self.assertIn("adjuntar el contrato firmado", pagina)
        self.assertNotIn('action="/auditor/requerimiento/generar"', pagina)

    def test_ruc_del_requerimiento_usa_el_de_la_auditoria(self):
        pagina = self.pagina(self.auditor)
        self.assertIn('name="ruc" value="0190377210001" readonly', pagina)
        with self.assertRaisesRegex(ValueError, "RUC del requerimiento debe coincidir"):
            self.confirmar(ruc="0190444619001")
        self.assertIsNone(self.contexto()["guardado"])
        with self.assertRaisesRegex(ValueError, "RUC del requerimiento debe coincidir"):
            R.save_requerimiento_datos(self.audit_id, normalizar_datos({**DATOS, "ruc": "0190444619001"}),
                                       self.auditor["id"], self.db)
        self.confirmar()
        self.assertEqual(self.contexto()["guardado"]["ruc"], self.audit["ruc"])

    def test_un_registro_antiguo_con_ruc_distinto_no_genera_documentos(self):
        self.contrato()
        self.confirmar()
        with connect(self.db) as conn:
            conn.execute("UPDATE requerimientos SET ruc = ? WHERE audit_id = ?", ("0190444619001", self.audit_id))
        with self.assertRaisesRegex(ValueError, "RUC del requerimiento no coincide"):
            self.generar()
        self.assertEqual(self.contexto()["paquetes"], [])
        pagina = self.pagina(self.auditor)
        self.assertIn("RUC del requerimiento: confirme los datos", pagina)


class TestAlmacenamientoArchivos(_Caso):
    def test_guardado_concurrente_no_reemplaza_ni_borra_el_archivo(self):
        with ThreadPoolExecutor(max_workers=8) as pool:
            resultados = list(pool.map(lambda _i: R._guardar(self.audit_id, PDF, "pdf", self.adjuntos), range(8)))
        self.assertEqual(sum(nuevo is not None for _ruta, _huella, nuevo in resultados), 1)
        self.assertEqual(R.ruta_archivo(resultados[0][0], self.adjuntos).read_bytes(), PDF)
        self.assertEqual(list(self.adjuntos.rglob("*.tmp")), [])


class TestPermisos(_Caso):
    def test_el_jefe_consulta_sin_formularios(self):
        self.listo()
        pagina = self.pagina(self.admin)
        self.assertIn("Modo solo lectura", pagina)
        self.assertNotIn('<form method="post"', pagina)
        self.assertIn("/requerimiento/archivo?", pagina, "Puede consultar las evidencias")

    def test_otro_auditor_no_ve_la_auditoria(self):
        self.assertIn("Auditoría no disponible", self.pagina(self.otro))

    def test_descargas_para_jefe_y_asignado_no_para_otro(self):
        self.listo()
        version = self.contexto()["versiones"]["carta"][0]
        query = {"audit_id": [str(self.audit_id)], "origen": ["generado"], "id": [str(version["id"])]}
        for user in (self.auditor, self.admin):
            contenido, nombre, tipo, _inline = router._descargar_archivo(user, query)
            self.assertTrue(contenido.startswith(b"%PDF"))
            self.assertEqual(tipo, "application/pdf")
        with self.assertRaises(PermissionError):
            router._descargar_archivo(self.otro, query)

    def test_un_archivo_de_otra_auditoria_no_se_descarga(self):
        self.listo()
        otra = create_company_audit("OTRA S.A.", "", "Quito", "", "2026", self.auditor["id"], self.admin["id"], self.db)
        version = self.contexto()["versiones"]["carta"][0]
        with self.assertRaisesRegex(ValueError, "no disponible"):
            router._descargar_archivo(self.auditor, {"audit_id": [str(otra)], "origen": ["generado"],
                                                     "id": [str(version["id"])]})

    def test_ruta_manipulada_no_sale_de_adjuntos(self):
        with self.assertRaisesRegex(ValueError, "no disponible"):
            R.ruta_archivo("../../auddit.db", self.adjuntos)

    def test_nombre_de_descarga_sin_comillas_ni_rutas(self):
        cabecera = content_disposition('CARTA "Ñ"/../x.pdf')
        self.assertTrue(cabecera.startswith('attachment; filename="CARTA ..x.pdf"'))
        self.assertIn("filename*=UTF-8''", cabecera)


class TestContratoYDatos(_Caso):
    def test_sin_contrato_no_se_genera(self):
        self.confirmar()
        with self.assertRaisesRegex(ValueError, "contrato"):
            self.generar()
        self.assertEqual(self.contexto()["versiones"]["carta"], [])

    def test_contrato_debe_ser_pdf_real(self):
        with self.assertRaisesRegex(ValueError, "no corresponde a un archivo PDF"):
            self.accion("/auditor/requerimiento/contrato",
                        _form(archivos={"archivo": [("contrato.pdf", b"MZ ejecutable")]}))
        with self.assertRaisesRegex(ValueError, "Formato no permitido"):
            self.accion("/auditor/requerimiento/contrato", _form(archivos={"archivo": [("contrato.exe", PDF)]}))

    def test_datos_incompletos_bloquean_la_generacion(self):
        self.contrato()
        mensaje = self.confirmar(representante_identificacion="", equipo="")
        self.assertIn("Pendiente", mensaje)
        with self.assertRaisesRegex(ValueError, "Cédula del representante.*Equipo de auditoría"):
            self.generar()

    def test_sin_confirmar_no_se_genera_con_la_precarga(self):
        self.contrato()
        with self.assertRaisesRegex(ValueError, "Confirme primero"):
            self.generar()

    def test_cedula_con_digito_verificador_invalido(self):
        with self.assertRaisesRegex(ValueError, "dígito verificador"):
            self.confirmar(representante_identificacion="0102030405")

    def test_precarga_solo_con_datos_respaldados(self):
        add_administrator(self.audit_id, CEDULA, "SERPA GARCIA EDUARDO", "ECUATORIANA", "GERENTE GENERAL",
                          self.db, user_id=self.auditor["id"])
        ctx = expedientes.get_audit_context(self.audit_id, self.db)
        sugeridos = precarga(self.audit, ctx)
        self.assertEqual(sugeridos["representante_identificacion"][0], CEDULA)
        self.assertEqual(sugeridos["anio_auditado"][0], "2026")
        self.assertNotIn("anio_certificados", sugeridos, "El año de los certificados no se deduce")
        self.assertNotIn("anio_cerrado", sugeridos, "Sin año fiscal confirmado no se propone")
        self.assertNotIn("fecha_corte", sugeridos)
        self.assertNotIn("representante_ciudad", sugeridos, "El domicilio no es la ciudad de la empresa")


class TestAnios(_Caso):
    def test_fecha_de_corte_dentro_del_anio_auditado(self):
        with self.assertRaisesRegex(ValueError, "año auditado"):
            normalizar_datos({**DATOS, "fecha_corte": "2025-07-31"})

    def test_ultimo_ejercicio_cerrado_no_posterior(self):
        with self.assertRaisesRegex(ValueError, "no puede ser posterior"):
            normalizar_datos({**DATOS, "anio_cerrado": "2027"})

    def test_formato_de_anio(self):
        with self.assertRaisesRegex(ValueError, "4 dígitos"):
            normalizar_datos({**DATOS, "anio_certificados": "25"})

    def test_cada_documento_usa_su_propio_anio(self):
        self.listo()
        ctx = self.contexto()
        carta = self._texto_pdf(ctx["versiones"]["carta"][0])
        certificado = self._texto_pdf(ctx["versiones"]["cert_paraisos"][0])
        self.assertIn("31 DE DICIEMBRE DE 2026", carta.replace("\n", " "))
        self.assertIn("en el año 2025", certificado.replace("\n", " "))
        self.assertNotIn("en el año 2026", certificado.replace("\n", " "))
        libro = self._libro(ctx["versiones"]["solicitud"][0])
        textos = [c.value for c in libro[HOJAS[1]]["B"] if c.value]
        self.assertTrue(any("del año 2025" in t for t in textos), "Req. #1 ítem 7: último ejercicio cerrado")
        self.assertTrue(any("durante el año 2026" in t for t in textos), "Req. #1 ítem 8: año auditado")
        textos2 = [c.value for c in libro[HOJAS[2]]["B"] if c.value]
        self.assertTrue(any("Enero a Julio de 2026" in t for t in textos2), "Req. #2: rango hasta la fecha de corte")

    def _texto_pdf(self, version) -> str:
        ruta = R.ruta_archivo(version["ruta"], self.adjuntos)
        return "\n".join(p.extract_text() for p in pypdf.PdfReader(ruta).pages)

    def _libro(self, version):
        return load_workbook(R.ruta_archivo(version["ruta"], self.adjuntos))


class TestGeneracionYVersiones(_Caso):
    def test_genera_cuatro_documentos_separados(self):
        self.listo()
        versiones = self.contexto()["versiones"]
        self.assertEqual({t for t, v in versiones.items() if v}, set(DOCUMENTOS))
        for tipo, (_n, extension, _p) in DOCUMENTOS.items():
            version = versiones[tipo][0]
            self.assertEqual(version["version"], 1)
            self.assertTrue(version["nombre"].endswith(f"v1.{extension}"))
            contenido = R.ruta_archivo(version["ruta"], self.adjuntos).read_bytes()
            self.assertEqual(len(version["sha256"]), 64)
            self.assertTrue(contenido.startswith(b"%PDF" if extension == "pdf" else b"PK"))
        self.assertEqual(self.contexto()["envios"], [], "Generar no registra ningún envío")

    def test_excel_de_cuatro_hojas_con_marcas_y_cuadros(self):
        self.contrato()
        self.confirmar()
        self.accion("/auditor/requerimiento/items", _form(item_1_2="cumplido", obs_1_3="ACTUALIZAR INFORMACIÓN",
                                                          item_2_19="no_aplica"))
        self.accion("/auditor/requerimiento/detalle", _form(seccion="bancos", filas=2, bancos_0_0="PRODUBANCO",
                                                            bancos_0_1="=HYPERLINK(\"http://x\")"))
        self.generar()
        ruta = R.ruta_archivo(self.contexto()["versiones"]["solicitud"][0]["ruta"], self.adjuntos)
        libro = load_workbook(ruta)
        self.assertEqual(libro.sheetnames, [HOJAS[1], HOJAS[2], HOJAS[3], HOJAS[4]])
        respuesta = leer_respuesta_xlsx(ruta.read_bytes())
        self.assertEqual(len(respuesta["items"]), len(ITEMS[1]) + len(ITEMS[2]))
        item2 = next(i for i in respuesta["items"] if (i["hoja"], i["numero"]) == (1, 2))
        self.assertTrue(item2["cumplido"])
        self.assertEqual(respuesta["detalles"]["bancos"][0][:2], ["PRODUBANCO", '=HYPERLINK("http://x")'])
        celda = next(c for fila in libro[HOJAS[4]].iter_rows() for c in fila if c.value == '=HYPERLINK("http://x")')
        self.assertEqual(celda.data_type, "s", "Un texto con = se guarda como texto, no como fórmula")

    def test_marcas_excluyentes(self):
        with self.assertRaisesRegex(ValueError, "estado no válido"):
            self.accion("/auditor/requerimiento/items", _form(item_1_2="ambos"))
        with self.assertRaisesRegex(ValueError, "no ambos"):
            R.save_requerimiento_items(self.audit_id, {(1, 2): {"cumplido": 1, "no_aplica": 1}}, self.db)

    def test_regenerar_crea_version_nueva_sin_tocar_la_anterior(self):
        self.listo()
        v1 = self.contexto()["versiones"]["carta"][0]
        archivo_v1 = R.ruta_archivo(v1["ruta"], self.adjuntos).read_bytes()
        self.confirmar(fecha_documentos="2026-09-05")
        self.generar()
        versiones = self.contexto()["versiones"]["carta"]
        self.assertEqual([v["version"] for v in versiones], [2, 1])
        self.assertNotEqual(versiones[0]["sha256"], v1["sha256"])
        self.assertEqual(R.ruta_archivo(versiones[1]["ruta"], self.adjuntos).read_bytes(), archivo_v1)
        self.assertEqual(versiones[1]["sha256"], v1["sha256"])

    def test_sin_cambios_no_se_crea_otra_version(self):
        self.listo()
        self.assertIn("Sin cambios", self.generar())
        self.assertEqual(len(self.contexto()["versiones"]["carta"]), 1)


def _ahora(minutos: int = 5) -> str:
    return (datetime.now() - timedelta(minutes=minutos)).strftime("%Y-%m-%dT%H:%M")


class _ConEnvio(_Caso):
    def _enviar(self, fecha: str | None = None, evidencia=("envio.png", PNG), paquete_id=None, **extra) -> str:
        archivos = {"evidencia": [evidencia]} if evidencia else {}
        paquete_id = paquete_id if paquete_id is not None else self.contexto()["paquetes"][0]["id"]
        return self.accion("/auditor/requerimiento/envio", _form(
            {"fecha": fecha or _ahora(0), "destinatario": "contadora@cliente.ec", "asunto": "INICIO",
             "paquete_id": paquete_id, **extra}, archivos))


class TestEnvios(_ConEnvio):
    def _versiones(self) -> dict:
        return {f"archivo_{t}": v[0]["id"] for t, v in self.contexto()["versiones"].items()}

    def test_registro_del_correo_con_evidencia(self):
        self.listo()
        with self.assertRaisesRegex(ValueError, "Seleccione el archivo"):
            self._enviar(evidencia=None)
        with self.assertRaisesRegex(ValueError, "posterior"):
            self._enviar(fecha="2099-01-01T10:00")
        self.assertIn("registrado", self._enviar())
        envio = self.contexto()["envios"][0]
        self.assertEqual(envio["canal"], "correo")
        self.assertEqual(sorted(json.loads(envio["archivos_json"])), sorted(self._versiones().values()))
        self.assertEqual(envio["paquete_id"], self.contexto()["paquetes"][0]["id"], "Queda qué generación se envió")
        self.assertIsNotNone(envio["evidencia_id"])

    def test_no_se_registra_una_mezcla_de_generaciones(self):
        self.listo()
        v1 = self._versiones()
        self.confirmar(fecha_documentos="2026-09-05")
        self.generar()
        v2 = self._versiones()
        mezcla = {**v2, "archivo_solicitud": v1["archivo_solicitud"]}
        # El formulario antiguo (una versión por documento) ya no basta: se indica la generación.
        with self.assertRaisesRegex(ValueError, "generación"):
            self.accion("/auditor/requerimiento/envio", _form(
                {"fecha": _ahora(), "destinatario": "a@b.ec", **mezcla}, {"evidencia": [("envio.png", PNG)]}))
        otra = create_company_audit("OTRA S.A.", "", "Quito", "", "2026", self.auditor["id"], self.admin["id"], self.db)
        with self.assertRaisesRegex(ValueError, "generación"):
            R.register_requerimiento_envio(otra, "correo", "a@b.ec", _ahora(), self.auditor["id"],
                                           paquete_id=self.contexto()["paquetes"][0]["id"], db_path=self.db)
        with self.assertRaisesRegex(ValueError, "generación"):
            self._enviar(paquete_id=99999)
        self.assertEqual(self.contexto()["envios"], [])
        self.assertEqual(self.contexto()["adjuntos"]["evidencia_correo"], [], "No se guarda evidencia de un envío inválido")

    def test_whatsapp_solo_despues_del_correo(self):
        self.listo()
        ahora = (datetime.now() - timedelta(minutes=1)).strftime("%Y-%m-%dT%H:%M")
        with self.assertRaisesRegex(ValueError, "primero el envío del correo"):
            self.accion("/auditor/requerimiento/whatsapp", _form(fecha=ahora, destinatario="Grupo"))
        self._enviar()
        self.assertIn("registrado", self.accion("/auditor/requerimiento/whatsapp", _form(fecha=ahora, destinatario="Grupo")))
        pagina = self.pagina(self.auditor)
        self.assertIn("Cliente notificado", pagina)
        self.assertIn("wa.me/?text=", pagina)

    def test_la_generacion_enviada_queda_identificada_tras_regenerar(self):
        self.listo()
        self._enviar()
        enviada = self.contexto()["paquetes"][0]
        archivo = R.ruta_archivo(enviada["archivos"]["carta"]["ruta"], self.adjuntos).read_bytes()
        self.confirmar(fecha_documentos="2026-09-07")
        self.generar()
        ctx = self.contexto()
        self.assertEqual([p["numero"] for p in ctx["paquetes"]], [2, 1])
        self.assertEqual(ctx["envios"][0]["paquete_id"], enviada["id"], "El envío histórico no cambia")
        self.assertEqual(R.ruta_archivo(ctx["paquetes"][1]["archivos"]["carta"]["ruta"], self.adjuntos).read_bytes(),
                         archivo)
        pagina = self.pagina(self.auditor)
        self.assertEqual(pagina.count("badge badge-blue\">Enviada"), 1)
        self.assertIn("Generación 2", pagina)
        self.assertIn('id="correo"', pagina)
        self.assertIn("Sin registrar", pagina.split('id="correo"', 1)[1].split("</section>", 1)[0])

    def test_borrador_eml_no_registra_envio(self):
        self.listo()
        contenido, _n, tipo, _i = router._borrador_correo(self.auditor, {"audit_id": [str(self.audit_id)]})
        self.assertEqual(tipo, "message/rfc822")
        self.assertIn(b"X-Unsent: 1", contenido)
        self.assertIn(b"X-Atlas-Generacion: 1 (REQ-", contenido)
        self.assertEqual(contenido.count(b"Content-Disposition: attachment"), 4)
        self.assertEqual(self.contexto()["envios"], [])


class _ConRespuesta(_ConEnvio):
    def _recibir(self, tipo: str, nombre: str, contenido: bytes) -> str:
        return self.accion("/auditor/requerimiento/recepcion", _form(
            {"tipo": tipo, "fecha": date.today().isoformat()}, {"archivo": [(nombre, contenido)]}))

    def _excel_generado(self, paquete=None) -> bytes:
        paquete = paquete or self.contexto()["paquetes"][0]
        return R.ruta_archivo(paquete["archivos"]["solicitud"]["ruta"], self.adjuntos).read_bytes()

    def _respondido(self, base: bytes | None = None, conflicto: bool = False, marca: str = "x",
                    cambio=None) -> bytes:
        """Excel generado por Atlas con la respuesta del cliente (ítem 4 del Req. #1)."""
        libro = load_workbook(io.BytesIO(base if base is not None else self._excel_generado()))
        for fila in libro[HOJAS[1]].iter_rows(min_row=9):
            if fila[0].value == 4:
                fila[2].value, fila[4].value = marca, "Se enviará el lunes; =no es fórmula"
                if conflicto:
                    fila[3].value = "X"
        if cambio:
            cambio(libro)
        salida = io.BytesIO()
        libro.save(salida)
        return salida.getvalue()

    def _enviado(self):
        self.listo()
        self._enviar()


class TestRecepcion(_ConRespuesta):
    def test_recepcion_parcial_y_pendientes(self):
        self._recibir("carta_firmada", "carta firmada.pdf", PDF)
        pagina = self.pagina(self.auditor)
        self.assertIn("Recepción incompleta", pagina)
        self.assertIn("Certificado de paraísos fiscales firmado", pagina)
        self.assertIn("0 de 4 completos", pagina, "Recibir no es revisar")
        self.assertIn("Recibido, pendiente de revisión", pagina)

    def test_excel_con_marcas_excluyentes_no_se_importa_pero_se_conserva(self):
        self._enviado()
        with self.assertRaisesRegex(ValueError, "conservado.*ítem 4: CUMPLIDO y NO APLICA"):
            self._recibir("solicitud_respondida", "respuesta.xlsx", self._respondido(conflicto=True))
        ctx = self.contexto()
        self.assertEqual(len(ctx["adjuntos"]["solicitud_respondida"]), 1, "El original se guarda")
        self.assertEqual(ctx["adjuntos"]["solicitud_respondida"][0]["importacion_estado"], "rechazado")
        self.assertEqual(ctx["respuestas"], [])
        self.assertIn("Recibido con error: no importado", self.pagina(self.auditor))

    def test_marca_no_reconocida(self):
        self._enviado()
        with self.assertRaisesRegex(ValueError, "marca no reconocida"):
            self._recibir("solicitud_respondida", "respuesta.xlsx", self._respondido(marca="tal vez"))

    def test_excel_valido_preserva_observaciones(self):
        self._enviado()
        mensaje = self._recibir("solicitud_respondida", "respuesta.xlsx", self._respondido())
        self.assertIn("1 de 37 ítems", mensaje)
        self.assertIn("generación 1", mensaje)
        respuesta = self.contexto()["respuestas"][0]
        items = json.loads(respuesta["items_json"])
        item4 = next(i for i in items if (i["hoja"], i["numero"]) == (1, 4))
        self.assertEqual(item4["observacion"], "Se enviará el lunes; =no es fórmula")
        self.assertTrue(item4["cumplido"] and not item4["no_aplica"])
        self.assertEqual(set(json.loads(respuesta["detalles_json"])), set(CUADROS))
        adjunto = self.contexto()["adjuntos"]["solicitud_respondida"][0]
        self.assertEqual(adjunto["importacion_estado"], "importado")
        self.assertEqual(adjunto["paquete_id"], self.contexto()["paquetes"][0]["id"])

    def test_excel_ajeno_o_danado(self):
        with self.assertRaisesRegex(ValueError, "no corresponde a un archivo XLSX"):
            self._recibir("solicitud_respondida", "respuesta.xlsx", PDF)
        with self.assertRaisesRegex(ValueError, "Falta la hoja"):
            leer_respuesta_xlsx(self._libro_vacio())
        with self.assertRaisesRegex(ValueError, "Formato no permitido"):
            self._recibir("carta_firmada", "carta.xlsx", b"PK\x03\x04")

    def _libro_vacio(self) -> bytes:
        from openpyxl import Workbook
        salida = io.BytesIO()
        Workbook().save(salida)
        return salida.getvalue()


# ── Hallazgos de la revisión técnica ────────────────────────────────────

class TestProcedenciaExcel(_ConRespuesta):
    """Un Excel respondido se importa solo si es de una generación enviada de
    esta auditoría: el nombre del archivo no cuenta."""

    def _rechazado(self, contenido: bytes, patron: str, estado: str = "rechazado") -> None:
        antes = len(self.contexto()["adjuntos"]["solicitud_respondida"])
        with self.assertRaisesRegex(ValueError, patron):
            self._recibir("solicitud_respondida", "REQUERIMIENTO CONSTRUCTORA ESGINGENIERIA.xlsx", contenido)
        ctx = self.contexto()
        self.assertEqual(len(ctx["adjuntos"]["solicitud_respondida"]), antes + 1, "Se conserva como evidencia")
        self.assertEqual(ctx["adjuntos"]["solicitud_respondida"][0]["importacion_estado"], estado)
        self.assertEqual(ctx["respuestas"], [], "No se importa")
        self.assertNotIn("recepcion", {p["clave"] for p in self._pasos_hechos()})

    def _pasos_hechos(self):
        from services.requerimiento import estado_proceso
        ctx = self.contexto()
        return [p for p in estado_proceso({
            "paquetes": ctx["paquetes"], "contratos": ctx["adjuntos"]["contrato"], "confirmado": True,
            "faltantes": [], "envios": ctx["envios"], "desactualizado": False,
            "recepciones": {t: ctx["adjuntos"][t] for t in ("carta_firmada", "cert_relacionadas_firmado",
                                                            "cert_paraisos_firmado", "solicitud_respondida")},
        })["pasos"] if p["hecho"]]

    def test_el_libro_tiene_cuatro_hojas_y_la_referencia_de_su_generacion(self):
        self._enviado()
        paquete = self.contexto()["paquetes"][0]
        libro = load_workbook(io.BytesIO(self._excel_generado()))
        self.assertEqual(libro.sheetnames, [HOJAS[1], HOJAS[2], HOJAS[3], HOJAS[4]], "Sin hoja de metadatos")
        propiedades = {p.name: p.value for p in libro.custom_doc_props}
        self.assertEqual(propiedades["AtlasReferencia"], paquete["referencia"])
        self.assertEqual(propiedades["AtlasRUC"], "0190377210001")
        self.assertEqual(propiedades["AtlasEjercicio"], "2026")
        self.assertIn(paquete["referencia"], libro[HOJAS[1]]["A2"].value)
        self.assertIn(paquete["referencia"], libro[HOJAS[2]]["A2"].value)

    def test_libro_valido_con_cualquier_nombre_de_archivo(self):
        self._enviado()
        mensaje = self.accion("/auditor/requerimiento/recepcion", _form(
            {"tipo": "solicitud_respondida", "fecha": date.today().isoformat()},
            {"archivo": [("OTRA EMPRESA 2019.xlsx", self._respondido())]}))
        self.assertIn("importado", mensaje)
        self.assertEqual(len(self.contexto()["respuestas"]), 1)

    def test_libro_de_otra_empresa(self):
        self._enviado()
        otra = create_company_audit("OTRA EMPRESA S.A.", "0190314014001", "Cuenca", "", "2026",
                                    self.auditor["id"], self.admin["id"], self.db)
        audit_otra = expedientes.get_audit(otra, self.auditor, self.db)
        contrato = R.save_requerimiento_adjunto(otra, "contrato", "c.pdf", PDF, "pdf", self.auditor["id"],
                                                db_path=self.db, adjuntos_dir=self.adjuntos)
        R.review_requerimiento_adjunto(otra, contrato, "conforme", "", self.auditor["id"], self.db)
        R.save_requerimiento_datos(otra, normalizar_datos({**DATOS, "empresa": "Otra Empresa S.A.",
                                                           "ruc": "0190314014001"}), self.auditor["id"], self.db)
        router.REQUERIMIENTO_POSTS["/auditor/requerimiento/generar"][1](_form(), audit_otra, self.auditor)
        paquete_otra = R.get_requerimiento_context(otra, self.db)["paquetes"][0]
        self._rechazado(self._respondido(self._excel_generado(paquete_otra)), "otra auditoría")

    def test_libro_con_otro_ruc(self):
        self._enviado()

        def otro_ruc(libro):
            libro.custom_doc_props["AtlasRUC"].value = "0190314014001"
        self._rechazado(self._respondido(cambio=otro_ruc), "ruc.*no coincide")

        def encabezado(libro):
            libro[HOJAS[1]]["A4"].value = libro[HOJAS[1]]["A4"].value.replace("0190377210001", "0190314014001")
        self._rechazado(self._respondido(cambio=encabezado), "encabezado")

    def test_libro_de_otro_ejercicio(self):
        self._enviado()

        def otro_anio(libro):
            libro.custom_doc_props["AtlasEjercicio"].value = "2025"
        self._rechazado(self._respondido(cambio=otro_anio), "ejercicio.*no coincide")

        def encabezado(libro):
            for hoja in (1, 2):
                libro[HOJAS[hoja]]["A4"].value = libro[HOJAS[hoja]]["A4"].value.replace("Ejercicio 2026",
                                                                                     "Ejercicio 2025")
        self._rechazado(self._respondido(cambio=encabezado), "encabezado")

    def test_paquete_incorrecto(self):
        self._enviado()
        g1 = self.contexto()["paquetes"][0]
        self.accion("/auditor/requerimiento/items", _form(item_1_2="cumplido"))
        self.generar()
        g2 = self.contexto()["paquetes"][0]

        # Libro de la generación 2 con la referencia de la 1: el contenido no coincide.
        def referencia_g1(libro):
            libro.custom_doc_props["AtlasReferencia"].value = g1["referencia"]
            for hoja in (1, 2):
                libro[HOJAS[hoja]]["A2"].value = libro[HOJAS[hoja]]["A2"].value.replace(g2["referencia"],
                                                                                     g1["referencia"])
        base_g2 = self._excel_generado(g2)
        self.confirmar(fecha_corte="2026-06-30")
        self.generar()
        g3 = self.contexto()["paquetes"][0]
        self._rechazado(self._respondido(self._excel_generado(g3), cambio=referencia_g1), "no coincide")

        # Referencias contradictorias entre la propiedad y la celda.
        def mezcla(libro):
            libro.custom_doc_props["AtlasReferencia"].value = g1["referencia"]
        self._rechazado(self._respondido(base_g2, cambio=mezcla), "no coinciden entre sí")

        # Referencia inventada.
        def inventada(libro):
            libro.custom_doc_props["AtlasReferencia"].value = "REQ-999-1-ABCDEF01"
            for hoja in (1, 2):
                libro[HOJAS[hoja]]["A2"].value = "Referencia Atlas: REQ-999-1-ABCDEF01"
        self._rechazado(self._respondido(base_g2, cambio=inventada), "ninguna generación")

        # Generación de esta auditoría que no consta como enviada: revisión manual, no importado.
        self._rechazado(self._respondido(base_g2), "revisión manual.*generación 2.*no consta como enviada",
                        estado="revision_manual")

    def test_libro_alterado(self):
        self._enviado()

        def detalle(libro):
            for fila in libro[HOJAS[2]].iter_rows(min_row=9):
                if fila[0].value == 5:
                    fila[1].value = "Otro texto"
        self._rechazado(self._respondido(cambio=detalle), "ítem 5.*no coincide")

        def quinta_hoja(libro):
            libro.create_sheet("Notas")
        self._rechazado(self._respondido(cambio=quinta_hoja), "exactamente las cuatro hojas")

        def celda_referencia(libro):
            libro[HOJAS[1]]["A2"].value = "Referencia Atlas: editada a mano"
        self._rechazado(self._respondido(cambio=celda_referencia), "alterado")

    def test_libro_antiguo_sin_identificacion(self):
        self._enviado()
        antiguo = self._respondido(build_solicitud_xlsx(normalizar_datos(DATOS), {}, {}))
        self._rechazado(antiguo, "revisión manual.*no trae la referencia", estado="revision_manual")
        ajeno = self._respondido(build_solicitud_xlsx(normalizar_datos({**DATOS, "ruc": "0190314014001"}), {}, {}))
        self._rechazado(ajeno, "no corresponde al RUC")
        self.assertIn("Requiere revisión manual: no importado", self.pagina(self.auditor))

    def test_vista_previa_no_se_puede_importar(self):
        self._enviado()
        previa, _n, _t, _i = router._vista_previa(self.auditor, {"audit_id": [str(self.audit_id)],
                                                                 "doc": ["solicitud"]})
        self.assertIn("VISTA PREVIA", load_workbook(io.BytesIO(previa))[HOJAS[1]]["A2"].value)
        self._rechazado(self._respondido(previa), "revisión manual", estado="revision_manual")

    def test_version_corregida_sin_borrar_el_original(self):
        self._enviado()
        with self.assertRaises(ValueError):
            self._recibir("solicitud_respondida", "respuesta.xlsx", self._respondido(conflicto=True))
        self.assertIn("importado", self._recibir("solicitud_respondida", "respuesta v2.xlsx", self._respondido()))
        recibidos = self.contexto()["adjuntos"]["solicitud_respondida"]
        self.assertEqual([r["importacion_estado"] for r in recibidos], ["importado", "rechazado"])
        for r in recibidos:
            self.assertTrue(R.ruta_archivo(r["ruta"], self.adjuntos).is_file())


class TestDesactualizacion(_ConEnvio):
    def _eml(self):
        from email import message_from_bytes, policy
        contenido, nombre, _t, _i = router._borrador_correo(self.auditor, {"audit_id": [str(self.audit_id)]})
        return message_from_bytes(contenido, policy=policy.default), nombre

    def test_editar_datos_marca_la_generacion_como_desactualizada(self):
        self.listo()
        v1 = self.contexto()["paquetes"][0]
        mensaje = self.confirmar(fecha_corte="2026-06-30")
        self.assertIn("generación 1 quedó desactualizada", mensaje)
        pagina = self.pagina(self.auditor)
        self.assertIn("Documentos desactualizados", pagina)
        self.assertNotIn("/requerimiento/correo.eml", pagina, "Sin borrador hasta regenerar")
        with self.assertRaisesRegex(ValueError, "cambiaron después de la generación 1"):
            self._eml()
        # La generación y sus archivos no cambian.
        self.assertEqual(self.contexto()["paquetes"][0]["datos_json"], v1["datos_json"])

    def test_cambiar_marcas_o_cuadros_tambien_desactualiza(self):
        self.listo()
        self.assertIn("desactualizada", self.accion("/auditor/requerimiento/items", _form(item_1_2="cumplido")))
        self.generar()
        self.assertIn("desactualizada", self.accion("/auditor/requerimiento/detalle",
                                                    _form(seccion="bancos", filas=1, bancos_0_0="PRODUBANCO")))

    def test_eml_coherente_con_su_generacion(self):
        self.listo()
        self.confirmar(fecha_corte="2026-06-30", correo_para="gerencia@cliente.ec")
        self.generar()
        paquete = self.contexto()["paquetes"][0]
        self.assertEqual(paquete["numero"], 2)
        eml, nombre = self._eml()
        self.assertIn("g2", nombre)
        self.assertEqual(eml["To"], "gerencia@cliente.ec")
        self.assertIn(paquete["referencia"], eml["X-Atlas-Generacion"])
        self.assertIn("30 de junio de 2026", eml.get_body(("plain",)).get_content())
        huellas = {hashlib.sha256(a.get_content()).hexdigest() for a in eml.iter_attachments()}
        self.assertEqual(huellas, {v["sha256"] for v in paquete["archivos"].values()})

    def test_la_pagina_muestra_el_correo_de_la_generacion(self):
        self.listo()
        self.confirmar(empresa="Nombre Nuevo S.A.")
        pagina = self.pagina(self.auditor)
        self.assertIn("CONSTRUCTORA ESGINGENIERIA S.A.S.", pagina.split('id="correo"', 1)[1].split("</section>")[0])

    def test_envio_de_una_generacion_anterior_queda_trazado(self):
        self.listo()
        g1 = self.contexto()["paquetes"][0]
        self.confirmar(fecha_corte="2026-06-30")
        with connect(self.db) as conn:
            conn.execute("UPDATE requerimiento_paquetes SET generado_at = ? WHERE id = ?",
                         ("2026-09-01 09:00:00", g1["id"]))
            conn.execute("UPDATE requerimientos SET updated_at = ? WHERE audit_id = ?",
                         ("2026-09-01 10:00:00", self.audit_id))
        # Un envío anterior se registra expresamente como histórico, con evidencia y motivo.
        self._enviar(fecha="2026-09-01T09:30", paquete_id=g1["id"], modo="historico",
                     justificacion_historica="Envío realizado antes de actualizar los datos")
        envio = self.contexto()["envios"][0]
        self.assertEqual(envio["paquete_id"], g1["id"])
        self.assertEqual(envio["registro_historico"], 1)
        self.assertIn("antes de actualizar", envio["justificacion_historica"])
        self.assertEqual(sorted(json.loads(envio["archivos_json"])), sorted(a["id"] for a in g1["archivos"].values()))

    def test_el_envio_nuevo_no_acepta_paquete_desactualizado(self):
        self.listo()
        paquete = self.contexto()["paquetes"][0]
        self.confirmar(fecha_corte="2026-06-30")
        with self.assertRaisesRegex(ValueError, "documentos desactualizados"):
            self._enviar(paquete_id=paquete["id"])
        with self.assertRaisesRegex(ValueError, "justificación|Explique"):
            self._enviar(paquete_id=paquete["id"], modo="historico")
        with self.assertRaisesRegex(ValueError, "antes de que esta generación"):
            self._enviar(paquete_id=paquete["id"], modo="historico",
                         justificacion_historica="Correo enviado antes de actualizar los datos")
        self.assertEqual(self.contexto()["envios"], [])
        pagina = self.pagina(self.auditor)
        self.assertNotIn('name="modo" value="actual"', pagina)
        self.assertIn('name="modo" value="historico"', pagina)

    def test_envio_historico_debe_ubicarse_antes_del_cambio(self):
        self.listo()
        paquete = self.contexto()["paquetes"][0]
        self.confirmar(fecha_corte="2026-06-30")
        with connect(self.db) as conn:
            conn.execute("UPDATE requerimiento_paquetes SET generado_at = ? WHERE id = ?",
                         ("2026-09-01 09:00:00", paquete["id"]))
            conn.execute("UPDATE requerimientos SET updated_at = ? WHERE audit_id = ?",
                         ("2026-09-01 10:00:00", self.audit_id))
        with self.assertRaisesRegex(ValueError, "antes de que esta generación"):
            self._enviar(paquete_id=paquete["id"], modo="historico",
                         justificacion_historica="Correo enviado después de editar datos")
        self.assertEqual(self.contexto()["envios"], [])

    def test_historial_de_generacion_anterior_respeta_el_cambio_previo(self):
        self.listo()
        anterior = self.contexto()["paquetes"][0]
        self.confirmar(fecha_corte="2026-06-30")
        self.generar()
        siguiente = self.contexto()["paquetes"][0]
        with connect(self.db) as conn:
            conn.execute("UPDATE requerimiento_paquetes SET generado_at = ? WHERE id = ?",
                         ("2026-09-01 09:00:00", anterior["id"]))
            conn.execute("UPDATE requerimiento_paquetes SET generado_at = ? WHERE id = ?",
                         ("2026-09-01 11:00:00", siguiente["id"]))
            conn.execute("UPDATE requerimientos SET updated_at = ? WHERE audit_id = ?",
                         ("2026-09-01 10:00:00", self.audit_id))
        with self.assertRaisesRegex(ValueError, "antes de que esta generación"):
            self._enviar(fecha="2026-09-01T10:30", paquete_id=anterior["id"], modo="historico",
                         justificacion_historica="Enviado después de cambiar los datos")
        self.assertEqual(self.contexto()["envios"], [])

    def test_cambio_solo_de_plantilla_permite_registro_historico(self):
        self.listo()
        with mock.patch("services.requerimiento.PLANTILLA_VERSION", "version-nueva"):
            self.assertIn("histórico", self._enviar(modo="historico",
                          justificacion_historica="Correo enviado con la plantilla anterior"))
        self.assertEqual(self.contexto()["envios"][0]["registro_historico"], 1)


class TestRevisionDeFirmados(_ConRespuesta):
    def test_contrato_adjuntado_no_revisado_no_permite_generar(self):
        self.assertIn("pendiente de revisión", self.contrato(revisar=False))
        self.confirmar()
        with self.assertRaisesRegex(ValueError, "revisión"):
            self.generar()
        self.assertEqual(self.contexto()["paquetes"], [])
        pagina = self.pagina(self.auditor)
        self.assertIn("Adjuntado, pendiente de revisión", pagina)
        self.assertIn("registrar la revisión del contrato firmado", pagina)
        self.assertIn("Atlas no valida criptográficamente", pagina)
        self.assertNotIn("Contrato revisado", pagina)
        self.assertIn("PDF legible, 1 página(s)", pagina)

    def test_contrato_revisado_con_constancia(self):
        self.contrato(revisar=False)
        contrato = self.contexto()["adjuntos"]["contrato"][0]
        with self.assertRaisesRegex(ValueError, "confirme cada verificación"):
            self.revisar(contrato["id"], todas=False)
        with self.assertRaisesRegex(ValueError, "qué se observó"):
            self.revisar(contrato["id"], resultado="observado", todas=False)
        self.assertIn("conforme", self.revisar(contrato["id"]))
        with self.assertRaisesRegex(ValueError, "ya fue revisado"):
            self.revisar(contrato["id"])
        revisado = self.contexto()["adjuntos"]["contrato"][0]
        self.assertEqual(revisado["revision_resultado"], "conforme")
        self.assertEqual(revisado["revisado_por"], self.auditor["id"])
        self.assertTrue(revisado["revisado_at"])
        self.confirmar()
        self.assertIn("generación 1", self.generar())
        pagina = self.pagina(self.auditor)
        self.assertIn("Revisado conforme", pagina)
        self.assertIn(f"Revisión de {self.auditor['full_name']}", pagina)

    def test_contrato_observado_no_habilita(self):
        self.contrato(revisar=False)
        self.revisar(self.contexto()["adjuntos"]["contrato"][0]["id"], resultado="observado",
                     nota="Falta la firma del gerente", todas=False)
        self.confirmar()
        with self.assertRaisesRegex(ValueError, "revisión"):
            self.generar()

    def test_pdf_ilegible_se_rechaza_aunque_tenga_cabecera(self):
        for contenido in (b"%PDF-1.4\n% solo cabecera\n", b"%PDF-1.7\n" + b"x" * 500):
            with self.assertRaisesRegex(ValueError, "No se pudo leer el PDF"):
                self.accion("/auditor/requerimiento/contrato",
                            _form(archivos={"archivo": [("contrato.pdf", contenido)]}))
        self.assertEqual(self.contexto()["adjuntos"]["contrato"], [])

    def test_solo_se_revisan_pdf_firmados_de_esta_auditoria(self):
        self._enviado()
        evidencia = self.contexto()["adjuntos"]["evidencia_correo"][0]
        with self.assertRaisesRegex(ValueError, "no requiere revisión"):
            self.revisar(evidencia["id"])

    def test_excel_invalido_no_completa_la_recepcion(self):
        self._enviado()
        for tipo in ("carta_firmada", "cert_relacionadas_firmado", "cert_paraisos_firmado"):
            self.assertIn("pendiente de revisión", self._recibir(tipo, f"{tipo}.pdf", PDF))
            self.revisar(self.contexto()["adjuntos"][tipo][0]["id"])
        with self.assertRaises(ValueError):
            self._recibir("solicitud_respondida", "respuesta.xlsx", self._respondido(conflicto=True))
        pagina = self.pagina(self.auditor)
        self.assertNotIn("Todo recibido", pagina)
        self.assertIn("Recibido con error", pagina)
        self.assertIn("3 de 4 completos", pagina)
        self.assertNotIn('req-step hecho actual', pagina)
        self._recibir("solicitud_respondida", "respuesta corregida.xlsx", self._respondido())
        self.assertIn("Todo recibido y revisado", self.pagina(self.auditor))

    def test_firmado_recibido_sin_revisar_no_cuenta(self):
        self._enviado()
        for tipo in ("carta_firmada", "cert_relacionadas_firmado", "cert_paraisos_firmado"):
            self._recibir(tipo, f"{tipo}.pdf", PDF)
        self._recibir("solicitud_respondida", "respuesta.xlsx", self._respondido())
        pagina = self.pagina(self.auditor)
        self.assertNotIn("Todo recibido", pagina)
        self.assertIn("1 de 4 completos", pagina)


class TestJefeSoloLectura(_Caso):
    def test_jefe_sin_vista_previa_ni_borrador_de_correo(self):
        self.listo()
        for funcion, extra in ((router._vista_previa, {"doc": ["carta"]}), (router._borrador_correo, {})):
            with self.assertRaisesRegex(PermissionError, "solo lectura"):
                funcion(self.admin, {"audit_id": [str(self.audit_id)], **extra})
        pagina = self.pagina(self.admin)
        self.assertNotIn("/requerimiento/vista-previa", pagina)
        self.assertNotIn("/requerimiento/correo.eml", pagina)
        self.assertNotIn("Revisar y dejar constancia", pagina)

    def test_jefe_descarga_versiones_y_evidencias_registradas(self):
        self.listo()
        ctx = self.contexto()
        for origen, fila in (("generado", ctx["paquetes"][0]["archivos"]["solicitud"]),
                             ("adjunto", ctx["adjuntos"]["contrato"][0])):
            contenido, _n, _t, _i = router._descargar_archivo(
                self.admin, {"audit_id": [str(self.audit_id)], "origen": [origen], "id": [str(fila["id"])]})
            self.assertEqual(hashlib.sha256(contenido).hexdigest(), fila["sha256"])

    def test_el_auditor_asignado_si_prepara(self):
        self.listo()
        contenido, _n, tipo, _i = router._vista_previa(self.auditor, {"audit_id": [str(self.audit_id)],
                                                                      "doc": ["carta"]})
        self.assertTrue(contenido.startswith(b"%PDF"))
        with self.assertRaises(PermissionError):
            router._vista_previa(self.otro, {"audit_id": [str(self.audit_id)], "doc": ["carta"]})


class TestGeneracionAtomica(_Caso):
    def _preparar(self):
        self.contrato()
        self.confirmar()

    def _sin_generacion(self):
        with connect(self.db) as conn:
            self.assertEqual(conn.execute("SELECT COUNT(*) FROM requerimiento_paquetes").fetchone()[0], 0)
            self.assertEqual(conn.execute("SELECT COUNT(*) FROM requerimiento_archivos").fetchone()[0], 0)
        carpeta = self.adjuntos / str(self.audit_id) / "requerimiento"
        restos = [f.name for f in carpeta.iterdir() if f.suffix == ".xlsx" or f.name.startswith(".")]
        self.assertEqual(restos, [], "No quedan archivos de la generación fallida")

    def test_fallo_al_guardar_el_tercer_documento(self):
        self._preparar()
        original, llamadas = R._guardar, []

        def falla_en_el_tercero(*args, **kwargs):
            llamadas.append(1)
            if len(llamadas) == 3:
                raise OSError("disco lleno (simulado)")
            return original(*args, **kwargs)
        with mock.patch.object(R, "_guardar", falla_en_el_tercero), self.assertRaises(OSError):
            self.generar()
        self._sin_generacion()
        pdfs = list((self.adjuntos / str(self.audit_id) / "requerimiento").glob("*.pdf"))
        self.assertEqual(len(pdfs), 1, "Solo queda el contrato; carta y certificado escritos se borraron")
        self.assertEqual(self.contexto()["paquetes"], [])

    def test_fallo_al_registrar_el_cuarto_documento(self):
        self._preparar()
        with connect(self.db) as conn:
            conn.execute("""CREATE TRIGGER falla_cuarto BEFORE INSERT ON requerimiento_archivos
                            WHEN NEW.tipo = 'solicitud' BEGIN SELECT RAISE(ABORT, 'fallo simulado'); END""")
        with self.assertRaisesRegex(Exception, "fallo simulado"):
            self.generar()
        self._sin_generacion()
        pagina = self.pagina(self.auditor)
        self.assertIn("Sin documentos generados", pagina)
        with connect(self.db) as conn:
            conn.execute("DROP TRIGGER falla_cuarto")
        self.assertIn("generación 1", self.generar())
        paquete = self.contexto()["paquetes"][0]
        self.assertEqual(set(paquete["archivos"]), set(DOCUMENTOS))
        self.assertEqual({v["version"] for v in paquete["archivos"].values()}, {1})

    def test_una_generacion_no_admite_documentos_sueltos(self):
        with self.assertRaisesRegex(ValueError, "cuatro documentos"):
            R.register_requerimiento_paquete(self.audit_id, 1, "REQ-1-1-AA", {"carta": ("c.pdf", PDF)}, {},
                                             self.auditor["id"], db_path=self.db, adjuntos_dir=self.adjuntos)


class TestMigracion(unittest.TestCase):
    """Una base con las tablas del requerimiento anteriores a esta revisión."""

    def test_migracion_idempotente_y_respuestas_previas_a_revision(self):
        db = Path(tempfile.mkdtemp()) / "vieja.db"
        with connect(db) as conn:
            conn.executescript("""
                CREATE TABLE requerimiento_archivos (id INTEGER PRIMARY KEY AUTOINCREMENT, audit_id INTEGER NOT NULL,
                    tipo TEXT NOT NULL, version INTEGER NOT NULL, nombre TEXT NOT NULL, ruta TEXT NOT NULL,
                    sha256 TEXT NOT NULL, bytes INTEGER NOT NULL, datos_json TEXT NOT NULL, generado_por INTEGER,
                    generado_at TEXT NOT NULL, UNIQUE (audit_id, tipo, version));
                CREATE TABLE requerimiento_adjuntos (id INTEGER PRIMARY KEY AUTOINCREMENT, audit_id INTEGER NOT NULL,
                    tipo TEXT NOT NULL, nombre TEXT NOT NULL, ruta TEXT NOT NULL, sha256 TEXT NOT NULL,
                    bytes INTEGER NOT NULL, fecha TEXT, nota TEXT NOT NULL DEFAULT '', subido_por INTEGER,
                    subido_at TEXT NOT NULL);
                CREATE TABLE requerimiento_envios (id INTEGER PRIMARY KEY AUTOINCREMENT, audit_id INTEGER NOT NULL,
                    canal TEXT NOT NULL, destinatario TEXT NOT NULL, copia TEXT NOT NULL DEFAULT '',
                    asunto TEXT NOT NULL DEFAULT '', fecha TEXT NOT NULL, archivos_json TEXT NOT NULL DEFAULT '[]',
                    evidencia_id INTEGER, registrado_por INTEGER, registrado_at TEXT NOT NULL);
                CREATE TABLE requerimiento_respuestas (id INTEGER PRIMARY KEY AUTOINCREMENT, audit_id INTEGER NOT NULL,
                    adjunto_id INTEGER NOT NULL, items_json TEXT NOT NULL, detalles_json TEXT NOT NULL,
                    importado_por INTEGER, importado_at TEXT NOT NULL);
                INSERT INTO requerimiento_adjuntos (audit_id, tipo, nombre, ruta, sha256, bytes, subido_at)
                    VALUES (1, 'solicitud_respondida', 'r.xlsx', 'x', 'y', 1, '2026-09-01'),
                           (1, 'contrato', 'c.pdf', 'x', 'y', 1, '2026-09-01');
                INSERT INTO requerimiento_respuestas (audit_id, adjunto_id, items_json, detalles_json, importado_at)
                    VALUES (1, 1, '[]', '{}', '2026-09-01');
            """)
        init_db(db, demo=True)
        init_db(db, demo=True)
        with connect(db) as conn:
            columnas = {r[1] for r in conn.execute("PRAGMA table_info(requerimiento_adjuntos)")}
            self.assertTrue({"revision_resultado", "revisado_por", "importacion_estado", "paquete_id"} <= columnas)
            self.assertIn("paquete_id", {r[1] for r in conn.execute("PRAGMA table_info(requerimiento_envios)")})
            filas = {r["tipo"]: r for r in conn.execute("SELECT * FROM requerimiento_adjuntos")}
        self.assertEqual(filas["solicitud_respondida"]["importacion_estado"], "revision_manual")
        self.assertIsNone(filas["contrato"]["revision_resultado"], "Un contrato previo no se da por revisado")
        self.assertEqual(R.get_requerimiento_context(1, db)["respuestas"], [])


class TestValidaciones(unittest.TestCase):
    def test_validar_archivo(self):
        self.assertEqual(validar_archivo("a.PDF", PDF, ("pdf",)), "pdf")
        self.assertEqual(validar_archivo("a.jpeg", b"\xff\xd8\xff\xe0", ("pdf", "jpg")), "jpg")
        with self.assertRaisesRegex(ValueError, "vacío"):
            validar_archivo("a.pdf", b"", ("pdf",))
        with self.assertRaisesRegex(ValueError, "10 MB"):
            validar_archivo("a.pdf", b"%PDF" + b"0" * (10 * 1024 * 1024), ("pdf",))
        with self.assertRaisesRegex(ValueError, "no parece un correo"):
            validar_archivo("a.eml", b"hola", ("eml",))

    def test_faltantes(self):
        datos = normalizar_datos(DATOS)
        self.assertEqual(faltantes(datos), [])
        self.assertIn("Fechas del cronograma de entrega de informes",
                      faltantes(normalizar_datos({**DATOS, "cronograma_2": ""})))
        with self.assertRaisesRegex(ValueError, "no es un correo válido"):
            normalizar_datos({**DATOS, "correo_cc": "fparra@accons"})


if __name__ == "__main__":
    unittest.main()
