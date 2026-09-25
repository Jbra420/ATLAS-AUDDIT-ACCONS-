"""
tests/test_resumen_excel.py — Levantamiento de información en Excel (/export/xlsx).

Genera el libro de un expediente y lo vuelve a abrir con openpyxl para
comprobar las hojas, los valores, los tipos de dato y la trazabilidad.
"""
from __future__ import annotations

import io
import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

from database import (
    add_administrator,
    add_shareholder,
    authenticate,
    connect,
    create_company_audit,
    get_audit,
    get_audit_context,
    init_db,
    mark_document_reviewed,
    patch_research,
    set_audit_fiscal_year,
    update_company_location_fields,
    update_company_profile_fields,
    upsert_financial_statement,
)
from services.resumen_excel import PENDIENTE, build_resumen_xlsx

RUC = "0190444619001"
HOJAS = [
    "Resumen", "1. SRI", "SRI - registro original", "2. Supercias", "Supercias - registro original",
    "3. Ubicación", "4. Administradores", "5. Accionistas", "6. Financiero", "EEFF - Situación",
    "EEFF - Resultados", "Requisitos", "Validaciones y alertas", "Documentos", "Fuentes y evidencia",
    "Hallazgos",
]


def _filas(ws) -> list[tuple]:
    return [tuple(v for v in fila) for fila in ws.iter_rows(values_only=True)]


def _valor(ws, etiqueta: str, columna: int = 2):
    """Valor de la fila cuya primera celda es la etiqueta."""
    for fila in ws.iter_rows():
        if fila[0].value == etiqueta:
            return fila[columna - 1]
    raise AssertionError(f"No se encontró {etiqueta!r} en {ws.title}")


class TestResumenExcel(unittest.TestCase):
    def setUp(self):
        self.db = Path(tempfile.mkdtemp()) / "test.db"
        init_db(self.db, demo=True)
        self.admin = authenticate("admin", "admin123", self.db)
        auditor = authenticate("auditor", "auditor123", self.db)
        self.uid = auditor["id"]
        self.audit_id = create_company_audit(
            "COBBLERCOMPANY CIA. LTDA.", RUC, "Cuenca", "", "2026", self.uid, self.admin["id"], self.db,
        )

    def _libro(self):
        audit = get_audit(self.audit_id, self.admin, self.db)
        return load_workbook(io.BytesIO(build_resumen_xlsx(audit, get_audit_context(self.audit_id, self.db))))

    def _completar(self):
        update_company_profile_fields(self.audit_id, {
            "razon_social_sri": "COBBLERCOMPANY CIA. LTDA.", "estado_contribuyente": "ACTIVO",
            "razon_social_supercias": "COBBLERCOMPANY CIA. LTDA.", "expediente_supercias": "712345",
        }, self.db, user_id=self.uid, fecha_consulta="2026-09-01")
        update_company_location_fields(self.audit_id, {"provincia": "AZUAY", "ciudad": "CUENCA"}, self.db,
                                       user_id=self.uid, fecha_consulta="2026-09-02")
        add_administrator(self.audit_id, "0102030405", "ORDOÑEZ FAJARDO JULIO", "ECUADOR", "GERENTE GENERAL",
                          db_path=self.db)
        add_shareholder(self.audit_id, 1, "0102030405", "ORDOÑEZ FAJARDO JULIO", self.db,
                        participacion_porcentaje="60", capital="600")
        add_shareholder(self.audit_id, 2, "1790012345001", "INVERSIONES ANDINAS S.A.", self.db,
                        participacion_porcentaje="40", capital="400")
        set_audit_fiscal_year(self.audit_id, 2025, user_id=self.uid, db_path=self.db)
        upsert_financial_statement(self.audit_id, 2025, {
            "activo_total": "1000", "pasivo_total": "600", "patrimonio_neto": "400",
            "ingresos_401": "900", "otros_ingresos_403": "100", "utilidad_neta_707": "50",
        }, fecha_consulta="2026-09-03", user_id=self.uid, db_path=self.db)

    def test_trae_una_hoja_por_bloque_y_las_de_control(self):
        self.assertEqual(self._libro().sheetnames, HOJAS)

    def test_expediente_vacio_marca_lo_pendiente(self):
        libro = self._libro()
        self.assertEqual(_valor(libro["1. SRI"], "Razón social").value, PENDIENTE)
        self.assertEqual(_valor(libro["1. SRI"], "RUC").value, RUC)
        self.assertIn(("Sin registros.",) + (None,) * 7, _filas(libro["4. Administradores"]))
        estados = [f[3] for f in _filas(libro["Requisitos"]) if f[2] == "Obligatorio"]
        self.assertIn("Pendiente", estados)

    def test_bloques_con_fuente_y_fecha_de_consulta(self):
        self._completar()
        libro = self._libro()
        razon = [c.value for c in libro["1. SRI"][_valor(libro["1. SRI"], "Razón social").row]]
        self.assertEqual(razon, ["Razón social", "COBBLERCOMPANY CIA. LTDA.", "SRI — Consulta de RUC",
                                 "2026-09-01", "Fuente registrada", "Coincide"])
        provincia = [c.value for c in libro["3. Ubicación"][_valor(libro["3. Ubicación"], "Provincia").row]]
        self.assertEqual(provincia[1:], ["AZUAY", "Supercias — Información general / Ubicación",
                                         "2026-09-02", "Fuente registrada", "Sin cruce automático"])
        self.assertIn("ORDOÑEZ FAJARDO JULIO", [f[2] for f in _filas(libro["4. Administradores"])])

    def test_accionistas_con_numeros_y_total(self):
        self._completar()
        ws = self._libro()["5. Accionistas"]
        capital, participacion = _valor(ws, 1, 5), _valor(ws, 1, 6)
        self.assertEqual(capital.value, 600.0)
        self.assertEqual(participacion.value, 0.6)
        self.assertEqual(participacion.number_format, "0.00%")
        total = _valor(ws, "Total", 5)
        self.assertEqual(total.value, f"=SUM(E{capital.row}:E{capital.row + 1})")

    def test_financiero_con_casilleros_e_indicadores(self):
        self._completar()
        ws = self._libro()["6. Financiero"]
        self.assertEqual(_valor(ws, "Año fiscal").value, 2025)
        self.assertEqual(_valor(ws, "Total activo (casillero 1)").value, 1000.0)
        self.assertEqual(_valor(ws, "Total ingresos (401 + 403) — calculado").value, 1000.0)
        endeudamiento = _valor(ws, "Razón de endeudamiento (Pasivo / Activo)")
        self.assertAlmostEqual(endeudamiento.value, 0.6)
        self.assertEqual(endeudamiento.number_format, "0.00%")

    def test_catalogos_originales_y_cuentas_completas_separados_del_expediente(self):
        self._completar()
        audit = get_audit(self.audit_id, self.admin, self.db)
        ctx = get_audit_context(self.audit_id, self.db)
        raw = build_resumen_xlsx(
            audit, ctx,
            catalog_sri={"ruc": RUC, "trade_name": "COBBLER", "parish": "SUCRE", "new_field": "dato"},
            catalog_supercias={"ruc": RUC, "representante": "PERSONA ORIGINAL", "_catalogo_fecha_actualizacion": "2026-09-20"},
            balance_details=[{
                "anio_fiscal": 2025, "archivo": "balances_2025_1.txt", "sha256": "a" * 64,
                "importado_at": "2026-09-24", "url": "https://example.com/origen",
                "cuentas": [("1", "ACTIVO", 1000.0), ("101", "CORRIENTE", 0.0),
                            ("2", "PASIVO", 600.0), ("3", "PATRIMONIO", 400.0),
                            ("401", "INGRESOS", 900.0), ("707", "GANANCIA", 50.0)],
            }],
        )
        libro = load_workbook(io.BytesIO(raw))
        sri = _filas(libro["SRI - registro original"])
        self.assertIn(("Nombre comercial", "trade_name", "COBBLER"), sri)
        self.assertIn(("Parroquia", "parish", "SUCRE"), sri)
        self.assertIn(("New field", "new_field", "dato"), sri)
        self.assertEqual(_valor(libro["Supercias - registro original"], "Representante legal", 3).value,
                         "PERSONA ORIGINAL")
        situacion = _filas(libro["EEFF - Situación"])
        resultados = _filas(libro["EEFF - Resultados"])
        self.assertIn((2025, "101", "CORRIENTE", 0.0, "Sin saldo"), situacion)
        self.assertIn((2025, "707", "GANANCIA", 50.0, "Con saldo"), resultados)
        self.assertNotIn((2025, "707", "GANANCIA", 50.0, "Con saldo"), situacion)
        self.assertIsNotNone(libro["EEFF - Situación"].freeze_panes)
        self.assertIn("A", libro["EEFF - Situación"].auto_filter.ref)

    def test_validaciones_requisitos_y_trazabilidad(self):
        self._completar()
        libro = self._libro()
        reglas = {f[0]: f[1] for f in _filas(libro["Validaciones y alertas"]) if f[0]}
        self.assertEqual(reglas["Razón social SRI = Supercias"], "Coincide")
        self.assertEqual(reglas["Activo = Pasivo + Patrimonio"], "Coincide")
        requisitos = {f[1]: f[3] for f in _filas(libro["Requisitos"]) if f[2] == "Obligatorio"}
        self.assertEqual(requisitos["Razón social (SRI)"], "Cumplido")
        self.assertEqual(requisitos["Gerente general registrado"], "Cumplido")
        self.assertEqual(requisitos["Presidente registrado"], "Pendiente")
        campos = [f[1] for f in _filas(libro["Fuentes y evidencia"])]
        self.assertIn("razon_social_sri", campos, "La trazabilidad lista cada dato registrado")

    def test_resumen_con_estado_y_recomendacion(self):
        self._completar()
        libro = self._libro()
        ws = libro["Resumen"]
        self.assertEqual(_valor(ws, "Empresa").value, "COBBLERCOMPANY CIA. LTDA.")
        self.assertEqual(_valor(ws, "Estado del RUC").value, "ACTIVO")
        self.assertEqual(_valor(ws, "Razón social (Supercias)").value, "COBBLERCOMPANY CIA. LTDA.")
        self.assertEqual(_valor(ws, "Total activo").value, 1000.0)
        self.assertEqual(_valor(ws, "Total activo (casillero 1)").value, 1000.0)
        self.assertEqual(_valor(ws, "Total ingresos (401 + 403)").value, 1000.0)
        self.assertEqual(_valor(ws, "Razón de endeudamiento").value, 0.6)
        self.assertEqual(_valor(ws, "Razón de endeudamiento").number_format, "0.00%")
        self.assertNotIn("Cobertura y soporte", [r[0] for r in _filas(ws)])
        self.assertIn("pendiente", _valor(ws, "Recomendación preliminar").value)
        self.assertIn("Objeto social", [r[1] for r in _filas(ws) if r[0] == "Obligatorio"])
        resumen = {r[:2] for r in _filas(ws) if isinstance(r[0], str)}
        for nombre in ("1. SRI", "2. Supercias", "3. Ubicación"):
            for campo in _filas(libro[nombre]):
                if campo[2] and campo[2] != "Fuente":
                    self.assertIn(campo[:2], resumen, f"Falta {campo[0]} de {nombre} en Resumen")

    def test_documentos_distinguen_pendiente_de_revision_registrada(self):
        doc = get_audit_context(self.audit_id, self.db)["docs"][0]
        mark_document_reviewed(self.audit_id, doc["id"], self.uid, self.db)
        libro = self._libro()
        self.assertEqual(_valor(libro["Documentos"], "Documentos revisados").value, "1 de 9")
        fila = next(r for r in _filas(libro["Documentos"]) if r[0] == doc["nombre"])
        self.assertEqual(fila[2], "Revisado")
        self.assertEqual(fila[3], "Auditor Demo")
        self.assertTrue(fila[4])
        self.assertEqual(sum(r[2] == "Pendiente" for r in _filas(libro["Documentos"])), 8)

    def test_recomendaciones_presentes_y_faltantes_figuran_en_la_matriz(self):
        update_company_profile_fields(self.audit_id, {"obligado_contabilidad": "SI"}, self.db,
                                      user_id=self.uid, fecha_consulta="2026-09-01")
        filas = _filas(self._libro()["Requisitos"])
        self.assertIn(("SRI", "Obligado a llevar contabilidad", "Recomendado", "Cumplido"), filas)
        self.assertIn(("Supercias", "Objeto social", "Obligatorio", "Pendiente"), filas)

    def test_detalle_financiero_por_cada_ejercicio_y_sin_total_parcial(self):
        self._completar()
        upsert_financial_statement(self.audit_id, 2024, {"activo_total": "300", "ingresos_401": "100"},
                                   fecha_consulta="2026-09-04", user_id=self.uid, db_path=self.db)
        ws = self._libro()["6. Financiero"]
        filas = _filas(ws)
        self.assertIn("Ejercicio 2025", [r[0] for r in filas])
        self.assertIn("Ejercicio 2024", [r[0] for r in filas])
        inicio_2024 = next(i for i, r in enumerate(filas) if r[0] == "Ejercicio 2024")
        bloque_2024 = filas[inicio_2024:]
        self.assertIn(("Total activo (casillero 1)", 300.0), [r[:2] for r in bloque_2024])
        self.assertIn(("Total ingresos (401 + 403) — calculado", PENDIENTE),
                      [r[:2] for r in bloque_2024])
        importe_2025 = next(r for r in filas if r[0] == "Total activo (casillero 1)")
        self.assertTrue(importe_2025[3].startswith("Supercias — Documentos económicos"))
        self.assertEqual(importe_2025[5], "Trazabilidad por dato")

    def test_accionistas_sin_montos_no_muestran_un_total_cero(self):
        add_shareholder(self.audit_id, 1, "0102030405", "SOCIA SIN MONTO", self.db)
        ws = self._libro()["5. Accionistas"]
        self.assertEqual(_valor(ws, "Total", 5).value, "Incompleto")
        self.assertEqual(_valor(ws, "Total", 6).value, "Incompleto")

    def test_hallazgos_incluyen_todas_las_notas_y_texto_largo(self):
        texto = "A" * 30010
        patch_research(self.audit_id, self.uid, {
            "supercias_info": "Expediente societario consultado", "sercop_info": "Sin registros",
            "pasted_text": texto,
        }, self.db)
        ws = self._libro()["Hallazgos"]
        self.assertEqual(_valor(ws, "Información Supercias").value, "Expediente societario consultado")
        self.assertEqual(_valor(ws, "Información SERCOP").value, "Sin registros")
        self.assertEqual(len(_valor(ws, "Texto de evidencia aportado (1/31)").value), 1000)
        self.assertEqual(len(_valor(ws, "Texto de evidencia aportado (31/31)").value), 10)

    def test_hojas_de_detalle_permiten_filtrar_y_mantienen_cabecera(self):
        libro = self._libro()
        for nombre in ("1. SRI", "4. Administradores", "Requisitos", "Documentos"):
            with self.subTest(hoja=nombre):
                self.assertIsNotNone(libro[nombre].freeze_panes)
                self.assertTrue(libro[nombre].auto_filter.ref)

    def test_texto_que_parece_formula_se_guarda_como_texto(self):
        patch_research(self.audit_id, self.uid, {"observations": '=HYPERLINK("http://x","clic")'}, self.db)
        celda = _valor(self._libro()["Hallazgos"], "Observaciones del auditor")
        self.assertEqual(celda.data_type, "s")
        self.assertEqual(celda.value, '=HYPERLINK("http://x","clic")')

    def test_no_modifica_el_expediente(self):
        """No genera ni guarda el resumen (a diferencia de la descarga en texto)."""
        get_audit_context(self.audit_id, self.db)  # como al abrir el expediente
        with connect(self.db) as conn:
            antes = conn.execute("SELECT generated_summary, updated_at FROM research_notes WHERE audit_id = ?",
                                 (self.audit_id,)).fetchone()
        self._libro()
        with connect(self.db) as conn:
            despues = conn.execute("SELECT generated_summary, updated_at FROM research_notes WHERE audit_id = ?",
                                   (self.audit_id,)).fetchone()
        self.assertEqual(tuple(antes or ()), tuple(despues or ()))


if __name__ == "__main__":
    unittest.main()
