"""Snapshot what users see and download, to prove a renaming changed nothing.

Check D of every naming phase. It never writes to the source database: the
source is copied with SQLite's backup API (opened read-only) into the output
folder, the app is pointed at that copy (ATLAS_DB_PATH), init_db migrates the
copy, and then every page and export is produced with the same functions the
server uses:

  pages/<name>.txt    visible text of the page (must stay identical)
  pages/<name>.html   full HTML, CSRF token blanked (diff is informative:
                      class and field names are expected to change)
  exports/<name>.txt  summary export and a cell-by-cell dump of the Excel
  documents/*.txt     engagement letter, certificates and information request
                      built from a fixed data set (text of each PDF, Excel cells)

Usage (from the repository root, before and after a phase):
  python3 scripts/naming/snapshot_outputs.py --source auddit.db --out /tmp/snap_before
  python3 scripts/naming/snapshot_outputs.py --source auddit.db --out /tmp/snap_after
  python3 scripts/naming/snapshot_outputs.py --compare /tmp/snap_before /tmp/snap_after
"""
from __future__ import annotations

import argparse
import difflib
import io
import os
import sqlite3
import sys
from html.parser import HTMLParser
from pathlib import Path

ROOT = Path(__file__).resolve().parents[2]
CSRF = "SNAPSHOT-CSRF"

# Fixed data for the generated documents (same shape as the confirmation form).
DOCUMENT_FORM = {
    "empresa": "Constructora Esgingenieria S.A.S.", "ruc": "0190377210001", "representante_titulo": "Mgtr.",
    "representante_nombre": "Eduardo Alfonso Serpa Garcia", "representante_cargo": "Gerente",
    "representante_identificacion": "0104926555", "representante_nacionalidad": "Ecuatoriana",
    "representante_ciudad": "Cuenca", "anio_auditado": "2026", "anio_certificados": "2025", "anio_cerrado": "2025",
    "fecha_documentos": "2026-09-01", "fecha_corte": "2026-07-31",
    "fechas_inventario": "entre el 15 de octubre y el 15 de diciembre",
    "auddit_representante": "Mgtr. Fernando Parra Suarez", "auddit_cargo": "Gerente",
    "correo_para_nombre": "Contadora", "correo_para": "contadora@cliente.ec", "correo_cc": "fparra@accons.ec",
    "equipo": "Mgtr. Fernando Parra Suarez\nLcda. Camila Guevara Lucero",
    **{f"cronograma_{i}": f"Hasta {m} 2027" for i, m in enumerate(("febrero", "marzo", "abril", "julio"))},
}


class _VisibleText(HTMLParser):
    """Text a person sees: text nodes and the values shown in form controls."""

    SKIP = {"script", "style", "head"}

    def __init__(self) -> None:
        super().__init__()
        self.parts: list[str] = []
        self._skip = 0

    def handle_starttag(self, tag, attrs):
        if tag in self.SKIP:
            self._skip += 1
        attrs = dict(attrs)
        if tag == "input" and attrs.get("type") not in ("hidden", "password", "file") and attrs.get("value"):
            self.parts.append(f"[{attrs['value']}]")
        if tag in ("br", "p", "div", "li", "tr", "h1", "h2", "h3", "section", "option", "label"):
            self.parts.append("\n")

    def handle_endtag(self, tag):
        if tag in self.SKIP and self._skip:
            self._skip -= 1

    def handle_data(self, data):
        if not self._skip and data.strip():
            self.parts.append(" ".join(data.split()))


def visible_text(html: str) -> str:
    parser = _VisibleText()
    parser.feed(html)
    lines = (" ".join(line.split()) for line in " ".join(parser.parts).split("\n"))
    return "\n".join(line for line in lines if line) + "\n"


def copy_database(source: Path, target: Path) -> None:
    if source.resolve() == target.resolve():
        raise SystemExit("The copy cannot be the source database")
    target.parent.mkdir(parents=True, exist_ok=True)
    if target.exists():
        target.unlink()
    with sqlite3.connect(f"file:{source}?mode=ro", uri=True) as src, sqlite3.connect(target) as dst:
        src.backup(dst)


def workbook_dump(content: bytes) -> str:
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(content))
    lines = [f"sheets: {wb.sheetnames}"]
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if cell.value not in (None, ""):
                    lines.append(f"{ws.title}!{cell.coordinate}\t{cell.value}")
    return "\n".join(lines) + "\n"


def pdf_text(content: bytes) -> str:
    from pypdf import PdfReader
    return "\n".join(page.extract_text() or "" for page in PdfReader(io.BytesIO(content)).pages) + "\n"


def snapshot(source: Path, out: Path) -> None:
    copy = out / "db" / "atlas.db"
    copy_database(source, copy)
    # Point the app at the copy before importing it (DB_PATH is read at import time).
    os.environ["ATLAS_DB_PATH"] = str(copy)
    os.environ["ATLAS_ADJUNTOS_DIR"] = str(out / "db" / "adjuntos")
    sys.path.insert(0, str(ROOT))
    os.chdir(ROOT)
    import database  # noqa: E402 — must follow the environment setup
    from core import router  # noqa: E402

    if Path(database.DB_PATH).resolve() != copy.resolve():
        raise SystemExit(f"Refusing to continue: the app points at {database.DB_PATH}")
    database.init_db(copy)

    with database.connect(copy) as conn:
        users = list(conn.execute("SELECT * FROM users WHERE active = 1 ORDER BY id"))
        audits = list(conn.execute("SELECT * FROM audits ORDER BY id"))
    by_id = {u["id"]: u for u in users}
    admin = next((u for u in users if u["role"] == "admin"), None)

    def write(kind: str, name: str, text: str) -> None:
        path = out / kind / name
        path.parent.mkdir(parents=True, exist_ok=True)
        path.write_text(text, encoding="utf-8")

    def page(name: str, path: str, user, query: dict | None = None) -> None:
        html = router.GET_ROUTES[path][2](user, query or {}, path, CSRF)
        write("pages", f"{name}.html", html.replace(CSRF, "CSRF"))
        write("pages", f"{name}.txt", visible_text(html))

    for user in users:
        if user["role"] == "admin":
            for path in ("/admin", "/admin/users", "/admin/companies"):
                page(f"{user['username']}{path.replace('/', '_')}", path, user)
        else:
            page(f"{user['username']}_auditor", "/auditor", user)
    for audit in audits:
        query = {"audit_id": [str(audit["id"])]}
        auditor = by_id.get(audit["assigned_auditor_id"])
        if auditor is not None:
            page(f"audit{audit['id']}_auditor_intake", "/auditor/radar", auditor, query)
            page(f"audit{audit['id']}_auditor_initial_request", "/auditor/requerimiento", auditor, query)
            row = database.get_audit(audit["id"], auditor)
            summary_build, xlsx_build = router.EXPORTS["/export/summary"][2], router.EXPORTS["/export/xlsx"][2]
            write("exports", f"audit{audit['id']}_summary.txt", summary_build(row))
            write("exports", f"audit{audit['id']}_intake_xlsx.txt", workbook_dump(xlsx_build(row)))
        if admin is not None:
            page(f"audit{audit['id']}_admin_intake", "/admin/audit", admin, query)
            page(f"audit{audit['id']}_admin_initial_request", "/admin/requerimiento", admin, query)

    from services.requerimiento import normalizar_datos
    from services.requerimiento_docs import build_carta, build_certificado, build_solicitud_xlsx
    data = normalizar_datos(DOCUMENT_FORM)
    write("documents", "engagement_letter.txt", pdf_text(build_carta(data)))
    for kind in ("cert_relacionadas", "cert_paraisos"):
        write("documents", f"{kind}.txt", pdf_text(build_certificado(kind, data)))
    write("documents", "information_request_xlsx.txt", workbook_dump(build_solicitud_xlsx(data, {}, {})))
    files = sorted(p.relative_to(out).as_posix() for p in out.rglob("*.txt"))
    print(f"{len(files)} snapshot files in {out} ({len(audits)} audits, {len(users)} active users)")


def compare(before: Path, after: Path) -> int:
    """Visible text must be identical; HTML differences are only reported."""
    failures = changed_html = 0
    names = sorted({p.relative_to(before).as_posix() for p in before.rglob("*") if p.is_file() and "db/" not in
                    p.relative_to(before).as_posix()})
    for name in names:
        a, b = before / name, after / name
        if not b.exists():
            print(f"MISSING  {name}")
            failures += 1
            continue
        if a.read_bytes() == b.read_bytes():
            continue
        if name.endswith(".html"):
            changed_html += 1
            continue
        failures += 1
        print(f"CHANGED  {name}")
        diff = difflib.unified_diff(a.read_text().splitlines(), b.read_text().splitlines(), "before", "after",
                                    lineterm="", n=1)
        for line in list(diff)[:20]:
            print("    " + line)
    extra = sorted({p.relative_to(after).as_posix() for p in after.rglob("*.txt")} -
                   {p.relative_to(before).as_posix() for p in before.rglob("*.txt")})
    for name in extra:
        print(f"NEW      {name}")
    print(f"{len(names)} files compared; HTML with markup changes: {changed_html}; failures: {failures}")
    print("OK: users see the same content" if not failures else "FAIL: visible content changed")
    return 1 if failures else 0


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--source", type=Path, help="database to copy (it is only read)")
    parser.add_argument("--out", type=Path, help="output folder for the snapshot")
    parser.add_argument("--compare", nargs=2, type=Path, metavar=("BEFORE", "AFTER"))
    args = parser.parse_args()
    if args.compare:
        return compare(*args.compare)
    if not (args.source and args.out):
        parser.error("use --source and --out, or --compare")
    if not args.source.is_file():
        parser.error(f"{args.source} does not exist")
    if args.out.resolve() == ROOT or ROOT in args.out.resolve().parents:
        parser.error("write snapshots outside the repository")
    snapshot(args.source, args.out)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
