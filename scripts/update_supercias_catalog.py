"""
scripts/update_supercias_catalog.py — Carga del Directorio de Compañías de
Supercías (catálogo local) para Atlas.

La Superintendencia de Compañías, Valores y Seguros publica un directorio
público, sin sesión ni CAPTCHA, descargable en formato .xlsx en:
  https://mercadodevalores.supercias.gob.ec/reportes/directorioCompanias.jsf

Este script descarga (o lee, si se le pasa un archivo local) ese directorio
y lo carga en una base local ultrarrápida (supercias_catalog.db) para que
Atlas la consulte instantáneamente por RUC, exactamente igual que
scripts/update_catastro.py hace con el catastro del SRI.

Nota técnica: el .xlsx que publica el portal trae un elemento <dimension>
incorrecto (declara "A1" aunque la hoja tiene ~227 mil filas), lo que hace
que openpyxl en modo read_only se detenga tras la primera fila. Antes de
abrirlo, este script reescribe ese único tag dentro del .xlsx (que sigue
siendo un ZIP) con un rango generoso, sin tocar el resto del archivo.

Uso:
  python scripts/update_supercias_catalog.py                    # descarga la última versión
  python scripts/update_supercias_catalog.py ruta/al/archivo.xlsx  # usa un archivo ya descargado
"""
from __future__ import annotations

import os
import re
import sqlite3
import sys
import tempfile
import unicodedata
import urllib.request
import zipfile
from pathlib import Path

DB_PATH = Path(__file__).parent.parent / "supercias_catalog.db"
DIRECTORY_URL = "https://mercadodevalores.supercias.gob.ec/reportes/excel/directorio_companias.xlsx"

# Encabezado oficial -> columna en supercias_catalog. Los que no aparecen
# aquí (No. FILA, REGIÓN, PRESENTÓ BALANCE INICIAL, FECHA PRESENTACIÓN
# BALANCE INICIAL) se ignoran a propósito: no forman parte del mapeo
# aprobado para el Radar Empresarial.
HEADER_MAP = {
    "EXPEDIENTE": "expediente",
    "RUC": "ruc",
    "NOMBRE": "razon_social",
    "SITUACION LEGAL": "situacion_legal",
    "FECHA_CONSTITUCION": "fecha_constitucion",
    "TIPO": "tipo_compania",
    "PAIS": "pais",
    "PROVINCIA": "provincia",
    "CANTON": "canton",
    "CIUDAD": "ciudad",
    "CALLE": "calle",
    "NUMERO": "numero",
    "INTERSECCION": "interseccion",
    "BARRIO": "barrio",
    "TELEFONO": "telefono",
    "REPRESENTANTE": "representante",
    "CARGO": "representante_cargo",
    "CAPITAL SUSCRITO": "capital_suscrito",
    "CIIU NIVEL 1": "ciiu_nivel1",
    "CIIU NIVEL 6": "ciiu_nivel6",
    "ULTIMO BALANCE": "ultimo_balance",
}
REQUIRED_COLUMNS = {"ruc", "razon_social"}
DB_COLUMNS = ["ruc", *sorted(c for c in HEADER_MAP.values() if c != "ruc")]


def _normalize_header(value: object) -> str:
    """Mayúsculas, sin acentos, espacios colapsados: para casar encabezados
    aunque el portal cambie tildes o mayúsculas/minúsculas entre versiones."""
    text = str(value or "").strip().upper()
    text = "".join(
        ch for ch in unicodedata.normalize("NFKD", text) if not unicodedata.combining(ch)
    )
    return re.sub(r"\s+", " ", text)


def _clean_date(value: str) -> str:
    value = (value or "").strip()
    parts = value.split("/")
    if len(parts) == 3 and all(p.isdigit() for p in parts):
        d, m, y = parts
        return f"{y}-{m.zfill(2)}-{d.zfill(2)}"
    return value


def _fetch_directory(dest: Path) -> None:
    print(f"Descargando directorio desde {DIRECTORY_URL} ...")
    request = urllib.request.Request(DIRECTORY_URL, headers={"User-Agent": "Atlas-Auddit/update_supercias_catalog"})
    with urllib.request.urlopen(request, timeout=180) as response, open(dest, "wb") as out:
        out.write(response.read())
    print(f"Descarga completa: {dest} ({dest.stat().st_size / 1_048_576:.1f} MB)")


def _patch_dimension(xlsx_path: Path) -> Path:
    """Corrige el <dimension> incorrecto del reporte y devuelve una copia
    temporal ya parcheada, lista para abrir con openpyxl en modo read_only.

    mkstemp() abre el archivo a nivel de sistema operativo (devuelve un file
    descriptor, no solo una ruta): hay que cerrarlo explícitamente apenas se
    obtiene la ruta, o el descriptor queda abierto durante toda la carga. Si
    la escritura del zip parchado falla a medio camino, se borra el temporal
    antes de propagar el error en vez de dejarlo huérfano.
    """
    fd, fixed_path_str = tempfile.mkstemp(suffix=".xlsx")
    os.close(fd)
    fixed_path = Path(fixed_path_str)
    try:
        with zipfile.ZipFile(xlsx_path, "r") as zin:
            sheet_bytes = zin.read("xl/worksheets/sheet1.xml")
            patched = re.sub(rb'<dimension ref="[^"]*"/>', b'<dimension ref="A1:AZ1000000"/>', sheet_bytes, count=1)

            with zipfile.ZipFile(fixed_path, "w", zipfile.ZIP_DEFLATED) as zout:
                for item in zin.infolist():
                    data = patched if item.filename == "xl/worksheets/sheet1.xml" else zin.read(item.filename)
                    zout.writestr(item, data)
    except Exception:
        fixed_path.unlink(missing_ok=True)
        raise
    return fixed_path


def _extract_metadata(rows_preview: list[tuple]) -> dict[str, str]:
    """Lee 'No. DE FILAS: N' y 'FECHA DE ACTUALIZACION: ...' de las primeras
    filas del reporte (son texto libre en la celda A, antes de la tabla)."""
    meta = {"total_filas": "", "fecha_actualizacion": ""}
    for row in rows_preview:
        first_cell = str(row[0] or "") if row else ""
        m = re.search(r"No\.\s*DE\s*FILAS\s*:\s*(\d+)", first_cell, re.I)
        if m:
            meta["total_filas"] = m.group(1)
        m = re.search(r"FECHA\s*DE\s*ACTUALIZACION\s*:\s*(.+)", first_cell, re.I)
        if m:
            meta["fecha_actualizacion"] = m.group(1).strip()
    return meta


def _init_db(conn: sqlite3.Connection) -> None:
    other_cols_sql = ", ".join(f"{c} TEXT" for c in DB_COLUMNS if c != "ruc")
    conn.execute(f"CREATE TABLE IF NOT EXISTS supercias_catalog (ruc TEXT PRIMARY KEY, {other_cols_sql})")
    conn.execute(
        """
        CREATE TABLE IF NOT EXISTS supercias_catalog_meta (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            total_filas TEXT,
            fecha_actualizacion TEXT,
            importado_at TEXT NOT NULL
        )
        """
    )
    conn.execute("CREATE INDEX IF NOT EXISTS idx_supercias_catalog_ruc ON supercias_catalog(ruc)")


def load_catalog(source_path: Path) -> None:
    import openpyxl  # import perezoso: solo esta herramienta CLI depende de openpyxl

    print(f"Procesando {source_path} ...")
    fixed_path = _patch_dimension(source_path)
    wb = None
    try:
        wb = openpyxl.load_workbook(fixed_path, read_only=True, data_only=True)
        ws = wb[wb.sheetnames[0]]
        row_iter = ws.iter_rows(values_only=True)

        preview_rows: list[tuple] = []
        header_row: tuple | None = None
        header_row_index = None
        for i, row in enumerate(row_iter):
            preview_rows.append(row)
            normalized = [_normalize_header(c) for c in row]
            if "RUC" in normalized and "NOMBRE" in normalized:
                header_row = row
                header_row_index = i
                break
            if i >= 20:
                break  # el encabezado real siempre está en las primeras filas

        if header_row is None:
            raise ValueError(
                "No se encontró la fila de encabezados (se esperaban columnas "
                "'RUC' y 'NOMBRE') en las primeras 20 filas del archivo. "
                "El formato del Directorio pudo haber cambiado; revisar manualmente."
            )

        meta = _extract_metadata(preview_rows[:header_row_index])

        col_index_by_field: dict[str, int] = {}
        for idx, raw_header in enumerate(header_row):
            field = HEADER_MAP.get(_normalize_header(raw_header))
            if field:
                col_index_by_field[field] = idx

        missing_required = REQUIRED_COLUMNS - set(col_index_by_field)
        if missing_required:
            raise ValueError(
                f"Columnas obligatorias ausentes en el encabezado: {sorted(missing_required)}. "
                "El formato del Directorio pudo haber cambiado; revisar manualmente antes de reintentar."
            )

        with sqlite3.connect(DB_PATH) as conn:
            conn.execute("PRAGMA synchronous = OFF")
            conn.execute("PRAGMA journal_mode = MEMORY")
            _init_db(conn)
            conn.execute("DELETE FROM supercias_catalog")

            placeholders = ", ".join("?" for _ in DB_COLUMNS)
            insert_sql = f"INSERT OR REPLACE INTO supercias_catalog ({', '.join(DB_COLUMNS)}) VALUES ({placeholders})"

            count = 0
            batch: list[tuple] = []
            for row in row_iter:
                ruc_idx = col_index_by_field.get("ruc")
                ruc = str(row[ruc_idx] or "").strip() if ruc_idx is not None else ""
                if not ruc or not ruc.isdigit():
                    continue
                nombre_idx = col_index_by_field.get("razon_social")
                nombre = str(row[nombre_idx] or "").strip() if nombre_idx is not None else ""
                if not nombre:
                    continue

                values = []
                for col in DB_COLUMNS:
                    idx = col_index_by_field.get(col)
                    raw = row[idx] if idx is not None and idx < len(row) else ""
                    text = str(raw if raw is not None else "").strip()
                    if col == "fecha_constitucion":
                        text = _clean_date(text)
                    values.append(text)
                batch.append(tuple(values))
                count += 1

                if len(batch) >= 10_000:
                    conn.executemany(insert_sql, batch)
                    batch.clear()
                    if count % 50_000 == 0:
                        print(f"{count} registros procesados...")

            if batch:
                conn.executemany(insert_sql, batch)

            conn.execute(
                "INSERT INTO supercias_catalog_meta (total_filas, fecha_actualizacion, importado_at) VALUES (?, ?, datetime('now', 'localtime'))",
                (meta["total_filas"] or str(count), meta["fecha_actualizacion"]),
            )

        print(f"Carga completa. {count} compañías importadas a {DB_PATH}.")
        print(f"Corte declarado por el portal: {meta['fecha_actualizacion'] or 'no declarado'} "
              f"({meta['total_filas'] or '?'} filas según el archivo).")
    finally:
        if wb is not None:
            wb.close()
        fixed_path.unlink(missing_ok=True)


def main() -> None:
    if len(sys.argv) >= 2:
        path = Path(sys.argv[1])
        if not path.exists():
            print(f"Error: el archivo {path} no existe.")
            sys.exit(1)
        load_catalog(path)
        return

    fd, tmp_path_str = tempfile.mkstemp(suffix=".xlsx")
    os.close(fd)
    tmp_path = Path(tmp_path_str)
    try:
        _fetch_directory(tmp_path)
        load_catalog(tmp_path)
    finally:
        tmp_path.unlink(missing_ok=True)


if __name__ == "__main__":
    main()
